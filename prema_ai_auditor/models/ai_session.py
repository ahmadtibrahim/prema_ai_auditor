# FILE: /opt/odoo/custum-addons/prema_ai_auditor/models/ai_session.py
"""
Prema AI Session — Full ERP access, web search, file generation,
voice transcription, CRM lead creation.
"""

import base64
import json
import logging
import math
import re
import traceback
import tempfile
import os

import requests
from dateutil.relativedelta import relativedelta
from odoo import api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

LEAD_GEN_DEFAULT_TITLES = [
    "Director of Logistics",
    "VP of Logistics",
    "VP Supply Chain",
    "Director of Supply Chain",
    "Supply Chain Manager",
    "Logistics Manager",
    "Transportation Manager",
    "Distribution Manager",
    "Shipping Manager",
    "Warehouse Manager",
    "Operations Manager",
    "Dispatch Manager",
    "Dispatch Supervisor",
    "Logistics Coordinator",
    "Transportation Coordinator",
    "Procurement Manager",
    "Purchasing Manager",
    "Fleet Manager",
]

GOOGLE_TEXT_SEARCH_FIELDS = ",".join([
    "places.id",
    "places.displayName",
    "places.formattedAddress",
    "places.location",
    "places.primaryType",
    "places.types",
])

GOOGLE_PLACE_DETAIL_FIELDS = ",".join([
    "id",
    "displayName",
    "formattedAddress",
    "location",
    "postalAddress",
    "websiteUri",
    "nationalPhoneNumber",
    "internationalPhoneNumber",
    "businessStatus",
    "primaryType",
    "types",
])


class PremaAISession(models.Model):
    _name = "prema.ai.session"
    _description = "Prema AI Session"

    name = fields.Char(string="Session Name", default="AI Chat")
    user_id = fields.Many2one(
        "res.users", default=lambda self: self.env.user, ondelete="cascade")
    message_ids = fields.One2many(
        "prema.ai.message", "session_id", string="Messages")
    last_attachment_id = fields.Many2one("prema.ai.attachment", string="Last Attachment")
    model_mode = fields.Selection([
        ("primary", "Primary"), ("fast", "Fast"), ("vision", "Vision"),
    ], default="primary", string="AI Model Mode")
    pending_lead_data = fields.Text(string="Pending Lead Data (JSON)")
    chat_mode = fields.Selection([
        ("standard", "Standard"),
        ("deep_thinking", "Deep Thinking"),
        ("web_research", "Web Research"),
        ("pricing_review", "Pricing Review"),
        ("logistics_estimate", "Logistics Estimate"),
        ("lead_generation", "Lead Generation"),
        ("document_qa", "Document Q&A"),
    ], default="standard", string="Chat Mode")

    # ── Session management ──────────────────────────────────────────
    @api.model
    def list_sessions(self):
        return self.search(
            [("user_id", "=", self.env.user.id)], order="create_date desc"
        ).read(["id", "name"])

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if "user_id" not in vals:
                vals["user_id"] = self.env.user.id
        return super().create(vals_list)

    @api.model
    def rename_session(self, session_id, new_name):
        s = self.browse(session_id)
        if not s.exists():
            raise UserError("Session not found")
        s.name = new_name
        return True

    @api.model
    def delete_session(self, session_id):
        s = self.browse(session_id).sudo()
        if not s.exists():
            return False
        s.unlink()
        return True

    @api.model
    def set_chat_mode(self, session_id, mode):
        """Persist the chat mode on the session without sending a message."""
        valid = {"standard", "deep_thinking", "web_research",
                 "pricing_review", "logistics_estimate", "lead_generation", "document_qa"}
        if mode not in valid:
            return False
        session = self.browse(session_id)
        if not session.exists():
            return False
        session.chat_mode = mode
        return True

    @api.model
    def get_active_trucks(self):
        """Return active fleet vehicles for the Logistics Estimate truck dropdown."""
        trucks = self.env["fleet.vehicle"].sudo().search([("active", "=", True)])
        return trucks.read([
            "id", "name", "license_plate",
            "x_reefer", "x_liftgate", "x_air_ride",
            "x_max_pallets", "x_max_payload_lbs", "x_gvwr_lbs",
            "x_vehicle_height_ft", "x_overall_length_ft",
            "x_tank_capacity_l", "x_current_fuel_percent", "x_estimated_range_km",
            "x_avg_km_per_l_last_week", "x_home_base_address", "x_last_location_address",
            "x_monthly_insurance_budget", "x_monthly_maintenance_budget",
        ])

    @api.model
    def get_pending_action(self, session_id):
        session = self.browse(session_id)
        if not session.exists():
            return None
        return session._build_pending_action_payload()

    def _load_pending_leads(self):
        self.ensure_one()
        if not self.pending_lead_data:
            return []
        try:
            leads = json.loads(self.pending_lead_data) or []
        except Exception:
            return []

        changed = False
        normalized = []
        for idx, lead in enumerate(leads):
            entry = self._prepare_pending_lead_entry(lead, idx)
            if entry != lead:
                changed = True
            normalized.append(entry)
        if changed:
            self.pending_lead_data = json.dumps(normalized)
        return normalized

    def _save_pending_leads(self, leads):
        self.ensure_one()
        prepared = [
            self._prepare_pending_lead_entry(lead, idx)
            for idx, lead in enumerate(leads or [])
        ]
        self.pending_lead_data = json.dumps(prepared) if prepared else False
        return prepared

    def _build_pending_action_payload(self):
        self.ensure_one()
        leads = self._load_pending_leads()
        if not leads:
            return None
        return {
            "type": "confirm_leads",
            "count": len(leads),
            "leads": [self._build_pending_preview_lead(lead) for lead in leads],
        }

    def _build_pending_preview_lead(self, lead):
        return {
            "pending_key": lead.get("pending_key"),
            "company_name": lead.get("company_name", "Unknown"),
            "contact_name": lead.get("contact_name", ""),
            "title": lead.get("title", ""),
            "city": lead.get("city", ""),
            "state": lead.get("state", ""),
            "phone": lead.get("phone", ""),
            "website": lead.get("website", ""),
            "email": lead.get("email", ""),
            "description": lead.get("description", ""),
            "status_note": lead.get("status_note", ""),
            "skip_create": bool(lead.get("skip_create")),
            "import_label": lead.get("import_label") or "Add to CRM",
        }

    def _prepare_pending_lead_entry(self, lead, idx=0):
        entry = dict(lead or {})
        company_name = (entry.get("company_name") or entry.get("name") or "").strip()
        contact_name = (entry.get("contact_name") or entry.get("name") or "").strip()
        email = (entry.get("email") or "").strip()
        website = (entry.get("website") or "").strip().rstrip("/")
        domain = self._normalize_domain(entry.get("domain") or website or self._extract_domain_from_email(email))
        country = entry.get("country") or entry.get("country_code") or ""
        state = entry.get("state") or entry.get("province") or entry.get("state_code") or ""
        description = (entry.get("description") or "").strip()

        entry.update({
            "company_name": company_name or "Unknown",
            "contact_name": contact_name if contact_name and contact_name != company_name else "",
            "title": (entry.get("title") or entry.get("job_title") or "").strip(),
            "email": email,
            "phone": (entry.get("phone") or "").strip(),
            "website": website,
            "domain": domain,
            "street": (entry.get("street") or "").strip(),
            "city": (entry.get("city") or "").strip(),
            "state": state.strip(),
            "zip": (entry.get("zip") or "").strip(),
            "country": country.strip(),
            "country_code": (entry.get("country_code") or "").strip(),
            "state_code": (entry.get("state_code") or "").strip(),
            "industry": (entry.get("industry") or "").strip(),
            "description": description,
            "pending_key": entry.get("pending_key") or entry.get("google_place_id") or f"pending-{idx + 1}",
        })

        preview = self._preview_crm_target(entry)
        entry.update(preview)
        return entry

    # ── Voice transcription via Whisper ─────────────────────────────
    @api.model
    def transcribe_audio(self, audio_b64, mimetype="audio/webm"):
        """Send audio to OpenAI Whisper API, return text."""
        api_key = self._get_api_key_static()
        if not api_key:
            return {"error": "API key not found. Set prema_ai.api_key in Settings."}
        try:
            audio_bytes = base64.b64decode(audio_b64)
            ext = "webm"
            if "mp4" in mimetype:
                ext = "mp4"
            elif "wav" in mimetype:
                ext = "wav"
            elif "ogg" in mimetype:
                ext = "ogg"

            with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
                tmp.write(audio_bytes)
                tmp_path = tmp.name

            try:
                with open(tmp_path, "rb") as f:
                    resp = requests.post(
                        "https://api.openai.com/v1/audio/transcriptions",
                        headers={"Authorization": f"Bearer {api_key}"},
                        files={"file": (f"audio.{ext}", f, mimetype)},
                        data={"model": "whisper-1"},
                        timeout=30,
                    )
                if resp.status_code != 200:
                    return {"error": f"Whisper API error ({resp.status_code}): {resp.text[:200]}"}
                return {"text": resp.json().get("text", "")}
            finally:
                os.unlink(tmp_path)
        except Exception as e:
            return {"error": f"Transcription failed: {e}"}

    @api.model
    def _get_api_key_static(self):
        p = self.env["ir.config_parameter"].sudo()
        return (p.get_param("prema_ai.api_key") or p.get_param("openai.api_key") or "").strip()

    # ── Messaging API ───────────────────────────────────────────────
    def send_message(self, message, attachment_data=None, chat_mode=None, truck_id=None):
        self.ensure_one()
        # Update session chat mode if provided
        if chat_mode:
            valid_modes = {"standard", "deep_thinking", "web_research",
                           "pricing_review", "logistics_estimate", "lead_generation", "document_qa"}
            if chat_mode in valid_modes:
                self.chat_mode = chat_mode
        effective_mode = self.chat_mode or "standard"

        msg = (message or "").strip()
        if not msg and not attachment_data:
            return ""

        # Build display message
        display = msg or ""
        if attachment_data:
            names = [a.get("filename", "file") for a in attachment_data]
            display = (msg + "\n" if msg else "") + "\n".join(
                [f"\U0001F4CE {n}" for n in names])

        self.env["prema.ai.message"].create({
            "session_id": self.id, "role": "user", "content": display})

        lower = (msg or "").lower()

        # ── Intercept bill creation ──────────────────────────────────
        create_cmds = (
            "create the draft", "create draft", "create bill",
            "create vendor bill", "make the draft", "make draft",
            "confirm create draft", "scan invoice", "extract invoice",
        )
        if any(c in lower for c in create_cmds):
            if attachment_data:
                reply = self._handle_bill_creation(attachment_data)
            elif self.last_attachment_id:
                result = self.env["prema.ai.attachment"].sudo().process_attachment_by_id(
                    self.last_attachment_id.id, confirmed=True)
                reply = self._format_process_result(result)
            else:
                reply = "No attachment found. Upload a file first."
            self.env["prema.ai.message"].create({
                "session_id": self.id, "role": "assistant", "content": reply})
            return self._wrap_reply(reply)

        # ── Intercept file generation ────────────────────────────────
        file_cmds = [
            "generate a report", "create a report", "make a report",
            "export to excel", "export to pdf", "export to word",
            "create excel", "create pdf", "create word", "make excel",
            "make pdf", "make word", "generate excel", "generate pdf",
            "download as", "save as", "create a file", "make a file",
        ]
        if any(c in lower for c in file_cmds):
            erp_context = self._gather_erp_context(msg, chat_mode=effective_mode)
            reply = self._handle_file_generation(msg, erp_context)
            self.env["prema.ai.message"].create({
                "session_id": self.id, "role": "assistant", "content": reply})
            return self._wrap_reply(reply)

        # ── Intercept CRM lead creation ──────────────────────────────
        lead_cmds = [
            "create lead", "add to crm", "add leads", "create leads",
            "save as lead", "add them to crm", "save leads", "save to crm",
            "add these to crm", "add these companies", "create these leads",
            "add companies to crm", "add to pipeline", "save to pipeline",
            "add as leads", "put in crm", "put them in crm",
            "add in crm", "add to the crm", "save to the crm",
            "import to crm", "send to crm", "move to crm",
            "create crm leads", "make crm leads", "add these leads",
            "add all to crm", "add all leads", "add as crm",
        ]
        # Also catch: any message with "crm" + an action verb
        _crm_actions = ["add", "create", "save", "put", "import", "send", "move", "make"]
        _lead_trigger = (
            any(c in lower for c in lead_cmds)
            or ("crm" in lower and any(w in lower for w in _crm_actions))
        )
        if _lead_trigger:
            reply = self._handle_lead_creation_from_chat(msg)
            self.env["prema.ai.message"].create({
                "session_id": self.id, "role": "assistant", "content": reply})
            return self._wrap_reply(reply)

        # ── Lead generation search intercepts ─────────────────────────
        _domain_hint = re.search(
            r'\b([\w-]+\.(?:com|ca|org|net|io|co|biz|us))\b', lower
        )

        # Industry/location search (no domain) → Google Places discovery.
        # (Snov.io direct domain search removed 2026-08-18 — subscription cancelled.)
        _search_intent = any(w in lower for w in [
            "find", "search", "get me", "look for", "give me",
            "discover", "generate", "pull", "fetch", "show me",
        ])
        _biz_words = any(w in lower for w in [
            "compan", "business", "lead", "prospect", "shipper", "broker",
            "supplier", "vendor", "distributor", "manufactur", "wholesal",
            "warehouse", "import", "export", "produce", "food", "grocery",
            "retailer", "logistics", "freight", "carrier", "3pl",
        ])
        _is_industry_search = (
            effective_mode == "lead_generation"
            and not _domain_hint
            and _search_intent
            and _biz_words
        )

        if _is_industry_search:
            reply = self._handle_industry_lead_search(msg)
            self.env["prema.ai.message"].create({
                "session_id": self.id, "role": "assistant", "content": reply})
            return self._wrap_reply(reply)

        # ── Tag creation via chat ──────────────────────────────────────
        _tag_triggers = ["create tag", "add tag", "make tag", "new tag", "create a tag"]
        if any(t in lower for t in _tag_triggers):
            reply = self._handle_tag_creation(msg)
            self.env["prema.ai.message"].create({
                "session_id": self.id, "role": "assistant", "content": reply})
            return self._wrap_reply(reply)

        # ── Normal AI call with ERP context ──────────────────────────
        # When user uploads files to analyze, skip the slow stored-attachment OCR scan
        _file_focus_words = {
            "analyze", "analyse", "read this", "what is this", "what does this",
            "summarize", "summarise", "explain this", "tell me about this",
            "extract", "what's in", "what is in", "review this", "look at this",
            "this document", "this file", "this attachment", "this invoice",
            "this pdf", "this image", "this photo", "this contract",
        }
        is_file_focus = bool(attachment_data) and any(k in lower for k in _file_focus_words)
        erp_context = self._gather_erp_context(
            msg, chat_mode=effective_mode, has_attachments=bool(attachment_data),
            skip_att_scan=is_file_focus)

        # ── Logistics Estimate: inject HERE routing + truck data ──────
        if effective_mode == "logistics_estimate":
            erp_context += self._gather_logistics_context(msg, truck_id=truck_id)

        file_contents = []
        if attachment_data:
            for att in attachment_data:
                file_contents.append({
                    "filename": att.get("filename", "file"),
                    "b64": att.get("file_b64", ""),
                    "mimetype": att.get("mimetype", "application/octet-stream"),
                })

        reply = self._call_openai(
            mode=self.model_mode or "primary",
            erp_context=erp_context,
            file_contents=file_contents,
            chat_mode=effective_mode,
        )
        self.env["prema.ai.message"].create({
            "session_id": self.id, "role": "assistant", "content": reply})
        return self._wrap_reply(reply)

    def _gather_logistics_context(self, user_msg, truck_id=None):
        """Call LogisticsEstimateService and return a formatted context string."""
        try:
            from odoo.addons.premafirm_ai_engine.services.logistics_estimate_service import (
                LogisticsEstimateService,
            )
            svc = LogisticsEstimateService(self.env)
            result = svc.run_estimate(user_msg=user_msg, truck_id=truck_id)
            if result:
                return "\n\nLOGISTICS ESTIMATE DATA:\n" + json.dumps(result, indent=2, default=str)
        except Exception as e:
            _logger.warning("LogisticsEstimateService error: %s", e)
            return f"\n\nLOGISTICS ESTIMATE SERVICE ERROR: {e}"
        return ""

    def _wrap_reply(self, reply):
        """Wrap reply string into structured dict, including any pending action."""
        return {
            "reply": reply,
            "pending_action": self._build_pending_action_payload(),
        }

    def _handle_bill_creation(self, attachment_data):
        att = attachment_data[0]
        try:
            result = self.env["prema.ai.attachment"].upload_and_extract(
                self.id, att["filename"], att["file_b64"], att["mimetype"])
            if result.get("error"):
                return f"\u274C Extraction failed: {result['error']}"
            self.last_attachment_id = result["attachment_id"]
            proc = self.env["prema.ai.attachment"].sudo().process_attachment_by_id(
                result["attachment_id"], confirmed=True)
            return self._format_process_result(proc)
        except Exception as e:
            return f"\u274C Bill creation error: {e}"

    def _format_process_result(self, result):
        if not isinstance(result, dict):
            return str(result)
        if result.get("success"):
            m = ("\u2705 Draft bill created.\n"
                 f"- Move ID: {result.get('move_id')}\n"
                 f"- Vendor: {result.get('partner')}\n"
                 f"- Amount: {result.get('amount')}")
            if result.get("duplicate_warning"):
                m += f"\n\u26A0 Duplicate: {result['duplicate_warning']}"
            return m
        if result.get("confirmation_required"):
            return "Confirmation required. Reply: 'confirm create draft'."
        if result.get("error"):
            return f"\u274C {result['error']}"
        return f"\u26A0 Unexpected: {result}"

    # ── File generation handler ─────────────────────────────────────
    def _handle_file_generation(self, user_msg, erp_context):
        lower = user_msg.lower()
        # Detect requested format
        fmt = "pdf"  # default
        if any(w in lower for w in ["excel", "xlsx", "spreadsheet"]):
            fmt = "xlsx"
        elif any(w in lower for w in ["word", "docx", "document"]):
            fmt = "docx"
        elif any(w in lower for w in ["csv"]):
            fmt = "csv"
        elif any(w in lower for w in ["powerpoint", "pptx", "presentation", "slides"]):
            return ("\u26A0 I cannot create PowerPoint files.\n"
                    "I can generate: PDF, Excel (.xlsx), Word (.docx), or CSV.\n"
                    "Which format would you like?")

        # Ask AI to generate structured data for the report
        report_prompt = (
            f"The user wants a {fmt.upper()} report. Their request: {user_msg}\n\n"
            f"Based on this ERP data, generate a JSON response with:\n"
            f"- 'title': report title\n"
            f"- 'headers': list of column headers\n"
            f"- 'rows': list of lists (data rows)\n"
            f"- 'summary': brief text summary\n"
            f"Return ONLY valid JSON, no markdown.\n\n"
            f"ERP DATA:\n{erp_context[:6000]}"
        )

        api_key = self._get_api_key()
        if not api_key:
            return "\u26A0 API key not found. Set prema_ai.api_key in Settings."

        model = self._resolve_model("primary")
        try:
            resp = requests.post(
                "https://api.openai.com/v1/responses",
                headers={"Authorization": f"Bearer {api_key}",
                         "Content-Type": "application/json"},
                json={
                    "model": model,
                    "instructions": "You generate structured report data as JSON. Return ONLY valid JSON.",
                    "input": [{"role": "user", "content": report_prompt}],
                    **( {"temperature": 0.1} if self._supports_temperature(model) else {} ),
                    "max_output_tokens": 4096,
                    "store": False,
                },
                timeout=90,
            )
            if resp.status_code != 200:
                return f"\u26A0 AI error generating report data: {resp.text[:300]}"

            data = resp.json()
            raw = data.get("output_text", "")
            if not raw:
                for item in data.get("output", []):
                    if item.get("type") == "message":
                        for block in item.get("content", []):
                            if block.get("type") == "output_text" and block.get("text"):
                                raw = block["text"].strip()
                                break
                    if raw:
                        break
            if not raw:
                return "AI returned empty report data."

            raw = re.sub(r"^```(?:json)?\s*", "", raw.strip())
            raw = re.sub(r"\s*```$", "", raw.strip())
            report_data = json.loads(raw)

        except json.JSONDecodeError:
            return "\u26A0 AI generated invalid report data. Please try again with a simpler request."
        except Exception as e:
            return f"\u26A0 Report generation error: {e}"

        # Generate the file
        try:
            from ..services.file_generator import generate_file
            file_bytes, filename = generate_file(fmt, report_data)

            # Save as ir.attachment
            att = self.env["ir.attachment"].sudo().create({
                "name": filename,
                "type": "binary",
                "datas": base64.b64encode(file_bytes).decode(),
                "res_model": "prema.ai.session",
                "res_id": self.id,
                "mimetype": {
                    "pdf": "application/pdf",
                    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "csv": "text/csv",
                }.get(fmt, "application/octet-stream"),
            })

            base_url = self.env["ir.config_parameter"].sudo().get_param("web.base.url")
            download_url = f"{base_url}/web/content/{att.id}?download=true"

            return (f"\u2705 Report generated: {filename}\n"
                    f"Download: {download_url}\n\n"
                    f"{report_data.get('summary', '')}")

        except Exception as e:
            _logger.error("File generation error: %s", traceback.format_exc())
            return f"\u274C File generation failed: {e}"

    # ── CRM lead creation from chat ─────────────────────────────────
    def _handle_lead_creation_from_chat(self, user_msg):
        """Extract leads from conversation and show a confirmation preview.
        Does NOT create any records — user must click 'Create Leads' to confirm.
        """
        if self.pending_lead_data:
            pending = self._build_pending_action_payload()
            if pending:
                return (
                    f"{pending['count']} lead(s) are already staged below.\n"
                    "Use the lead cards to add one at a time, or click Create Leads to import the remaining results."
                )

        # Gather context from recent assistant messages
        recent = self.message_ids.filtered(lambda m: m.role == "assistant")[-5:]
        context = "\n".join([m.content for m in recent])

        api_key = self._get_api_key()
        if not api_key:
            return "API key not found. Set prema_ai.api_key in Settings."

        model = self._resolve_model("primary")
        prompt = (
            "From the following conversation, extract a list of business leads.\n"
            "Return ONLY a valid JSON array. Each lead object must include:\n"
            "- company_name (required — use 'Unknown' if not available)\n"
            "- contact_name (person name if mentioned)\n"
            "- website (if available)\n"
            "- street, city, state, zip, country\n"
            "- phone (if available)\n"
            "- email (if available)\n"
            "- industry (if determinable)\n"
            "- description (brief reason this is a lead)\n\n"
            f"CONVERSATION:\n{context}\n\nUser request: {user_msg}"
        )

        try:
            resp = requests.post(
                "https://api.openai.com/v1/responses",
                headers={"Authorization": f"Bearer {api_key}",
                         "Content-Type": "application/json"},
                json={
                    "model": model,
                    "instructions": "Extract business leads as JSON array. Return ONLY valid JSON, no markdown.",
                    "input": [{"role": "user", "content": prompt}],
                    **( {"temperature": 0.1} if self._supports_temperature(model) else {} ),
                    "max_output_tokens": 4096,
                    "store": False,
                },
                timeout=60,
            )
            if resp.status_code != 200:
                return f"AI error: {resp.text[:200]}"

            data = resp.json()
            # Responses API may put text in output_text OR output[].content[].text
            raw = data.get("output_text", "")
            if not raw:
                for item in data.get("output", []):
                    if item.get("type") == "message":
                        for block in item.get("content", []):
                            if block.get("type") == "output_text" and block.get("text"):
                                raw = block["text"].strip()
                                break
                    if raw:
                        break

            _logger.info("Lead extraction response (first 400): %s", raw[:400] if raw else "EMPTY")

            if not raw:
                return (
                    "AI returned an empty response. Please try again."
                )

            # Strip markdown code fences if present
            raw = re.sub(r"^```(?:json)?\s*", "", raw.strip())
            raw = re.sub(r"\s*```$", "", raw.strip())

            # If the text starts before '[', find the JSON array
            bracket = raw.find("[")
            if bracket > 0:
                raw = raw[bracket:]
            end_bracket = raw.rfind("]")
            if end_bracket != -1:
                raw = raw[:end_bracket + 1]

            leads = json.loads(raw)

            if not isinstance(leads, list) or not leads:
                return (
                    "No leads found in the conversation.\n"
                    "Try researching companies first, then say 'add to CRM'."
                )

            # Store pending leads — NOT created yet
            self._save_pending_leads(leads)

            # Build plain-text preview
            lines = [f"{len(leads)} lead(s) ready to add to CRM:\n"]
            for i, ld in enumerate(leads, 1):
                name = ld.get("company_name", "Unknown")
                parts = [f"{i}. {name}"]
                loc = ", ".join(filter(None, [ld.get("city"), ld.get("country")]))
                if loc:
                    parts.append(f"   Location: {loc}")
                if ld.get("phone"):
                    parts.append(f"   Phone: {ld['phone']}")
                if ld.get("website"):
                    parts.append(f"   Website: {ld['website']}")
                if ld.get("email"):
                    parts.append(f"   Email: {ld['email']}")
                if ld.get("description"):
                    parts.append(f"   Note: {ld['description'][:120]}")
                lines.append("\n".join(parts))

            lines.append(
                "\nReview the leads above, then click [Create Leads] to add them to CRM, "
                "or [Cancel] to discard."
            )
            return "\n\n".join(lines)

        except json.JSONDecodeError as jde:
            _logger.warning("Lead JSON parse error: %s | raw was: %s", jde, raw[:300] if 'raw' in dir() else "N/A")
            return (
                "Could not parse lead data from AI response.\n"
                "Try researching companies first, then say 'add to CRM'."
            )
        except Exception as e:
            _logger.error("Lead extraction error: %s", e)
            return f"Lead extraction error: {e}"

    # ── Industry-based lead search (Google Places → ranked leads) ─────
    def _handle_industry_lead_search(self, user_msg):
        """
        Lead-generation pipeline:
        1. Parse the request into location/radius/target sectors.
        2. Search Google Places for real companies in that area.
        3. Rank the candidates and stage them for CRM import.

        (Snov.io contact enrichment removed 2026-08-18 — subscription cancelled.)
        """
        if not self._get_google_maps_api_key():
            return (
                "Google Maps / Places API key not found. Set google_maps_api_key first, "
                "then try the lead search again."
            )

        try:
            plan = self._plan_lead_generation_request(user_msg)
        except UserError as exc:
            return str(exc)

        if not plan.get("location"):
            return (
                "I need a target city or area for lead generation. Example:\n"
                "\"Generate 20 leads in Ottawa, ON within 100 km for food distribution\""
            )

        try:
            candidates = self._prepare_google_lead_candidates(plan)
        except Exception as exc:
            _logger.warning("Lead generation search failed: %s", exc)
            return f"Lead search failed: {exc}"
        if not candidates:
            return (
                f"I could not find matching companies in Google Places for {plan['location']}.\n"
                "Try widening the radius, adding another industry term, or naming a nearby city."
            )

        ranked = self._rank_lead_candidates(user_msg, candidates, plan)
        target_count = max(1, min(int(plan.get("count") or 10), 50))
        final_leads = ranked[:target_count]
        staged = self._save_pending_leads(final_leads)
        with_contacts = sum(1 for lead in staged if lead.get("contact_name") or lead.get("email"))
        update_count = sum(1 for lead in staged if lead.get("existing_lead_id"))
        reuse_count = sum(1 for lead in staged if lead.get("existing_company_id") and not lead.get("existing_lead_id"))
        already_in_crm = sum(1 for lead in staged if lead.get("skip_create"))

        lines = [
            (
                f"Prepared {len(staged)} lead(s) for **{plan['location']}** within "
                f"**{int(plan['radius_km'])} km** using **Google Places** for company discovery."
            ),
            (
                f"{with_contacts} lead(s) include a named contact. "
                f"{update_count} will update an existing CRM lead. "
                f"{reuse_count} will reuse an existing company record."
            ),
        ]
        if already_in_crm:
            lines.append(
                f"{already_in_crm} result(s) already exist in CRM and are marked accordingly on the cards below."
            )
        lines.append(
            "Use the lead cards below to add one result at a time, or click Create Leads to import all remaining results."
        )
        return "\n\n".join(lines)

    def _get_google_maps_api_key(self):
        params = self.env["ir.config_parameter"].sudo()
        return (
            params.get_param("google_maps_api_key")
            or params.get_param("google.maps.api.key")
            or ""
        ).strip()

    def _resolve_lead_gen_model(self):
        params = self.env["ir.config_parameter"].sudo()
        return (
            params.get_param("prema_ai.lead_gen_model")
            or params.get_param("prema_ai.fast_model")
            or params.get_param("prema_ai.primary_model")
            or "gpt-5.5"
        )

    def _plan_lead_generation_request(self, user_msg):
        plan = self._default_lead_generation_plan(user_msg)
        api_key = self._get_api_key()
        if not api_key:
            return plan

        titles = self._get_lead_gen_titles()
        prompt = (
            f"User request: {user_msg}\n\n"
            "Extract the lead-generation search plan as JSON with these keys:\n"
            "- count: integer number of target leads (default 10)\n"
            "- location: city/region text like 'Ottawa, ON'\n"
            "- radius_km: integer search radius in kilometers (default 50)\n"
            "- search_terms: array of 2-6 Google Places search terms focused on shippers, distributors, "
            "manufacturers, wholesalers, importers, exporters, warehouses, produce and food logistics businesses\n"
            "- desired_titles: array of contact job titles to look for\n\n"
            f"Use these fallback titles if the user did not specify titles: {titles}\n"
            "Return ONLY valid JSON."
        )
        try:
            raw = self._call_responses_text(
                instructions="Convert lead generation requests into JSON search plans. Return only JSON.",
                prompt=prompt,
                model=self._resolve_lead_gen_model(),
                max_output_tokens=1200,
                temperature=0.1,
                timeout=45,
            )
            parsed = self._parse_json_payload(raw, expected="object") or {}
            if isinstance(parsed, dict):
                if parsed.get("count"):
                    plan["count"] = parsed.get("count")
                if parsed.get("location"):
                    plan["location"] = parsed.get("location")
                if parsed.get("radius_km"):
                    plan["radius_km"] = parsed.get("radius_km")
                if parsed.get("search_terms"):
                    plan["search_terms"] = parsed.get("search_terms")
                if parsed.get("desired_titles"):
                    plan["desired_titles"] = parsed.get("desired_titles")
        except Exception as exc:
            _logger.warning("Lead generation planning fallback: %s", exc)

        plan["count"] = max(1, min(int(plan.get("count") or 10), 50))
        plan["radius_km"] = max(5, min(int(plan.get("radius_km") or 50), 250))
        plan["search_terms"] = [
            str(term).strip() for term in (plan.get("search_terms") or []) if str(term).strip()
        ][:6]
        if not plan["search_terms"]:
            plan["search_terms"] = ["freight shipper", "distribution center", "warehouse logistics"]
        plan["desired_titles"] = [
            str(title).strip() for title in (plan.get("desired_titles") or self._get_lead_gen_titles())
            if str(title).strip()
        ][:12]
        return plan

    def _default_lead_generation_plan(self, user_msg):
        lower = (user_msg or "").lower()
        count_match = re.search(r"(\d+)\s+(?:lead|leads|company|companies|prospect|prospects)", lower)
        radius_match = re.search(r"(\d+(?:\.\d+)?)\s*(km|kilometers?|kilometres?|mi|miles?)", lower)

        count = int(count_match.group(1)) if count_match else 10
        radius_km = 50
        if radius_match:
            value = float(radius_match.group(1))
            unit = radius_match.group(2)
            radius_km = int(round(value * 1.60934)) if unit.startswith("mi") else int(round(value))

        parts = [part.strip() for part in (user_msg or "").split(",") if part.strip()]
        location = ""
        search_terms = []
        radius_idx = next(
            (idx for idx, part in enumerate(parts) if re.search(r"\d+(?:\.\d+)?\s*(km|mi)", part.lower())),
            None,
        )
        if radius_idx is not None:
            before = parts[:radius_idx]
            after = parts[radius_idx + 1:]
            if before:
                head = re.sub(
                    r"^(generate|find|search|show|give|get)\s+(me\s+)?\d+\s+(lead|leads|company|companies|prospect|prospects)\s*",
                    "",
                    before[0],
                    flags=re.IGNORECASE,
                ).strip()
                location_parts = [head] + before[1:]
                location = ", ".join([part for part in location_parts if part])
            search_terms = after

        if not location:
            near_match = re.search(
                r"(?:in|around|near|within)\s+([A-Za-z0-9 .'-]+(?:,\s*[A-Za-z]{2,})?)",
                user_msg or "",
                re.IGNORECASE,
            )
            if near_match:
                location = near_match.group(1).strip()

        if not search_terms:
            trailing = re.split(r"\b(?:radius|km|mile|miles|kilometers?|kilometres?)\b", user_msg or "", maxsplit=1)
            if len(trailing) > 1:
                search_terms = [part.strip() for part in trailing[-1].split(",") if part.strip()]

        if not search_terms:
            keyword_chunks = [
                chunk.strip()
                for chunk in re.split(r",| and | / ", user_msg or "", flags=re.IGNORECASE)
                if chunk.strip()
            ]
            search_terms = keyword_chunks[-3:]

        clean_terms = []
        for term in search_terms:
            cleaned = re.sub(
                r"\b(generate|find|get me|show me|give me|look for|lead|leads|radius|within|\d+\s*(km|mi))\b",
                "",
                term,
                flags=re.IGNORECASE,
            ).strip(" -")
            if cleaned:
                clean_terms.append(cleaned)

        return {
            "count": count,
            "location": location,
            "radius_km": radius_km,
            "search_terms": clean_terms[:6],
            "desired_titles": self._get_lead_gen_titles(),
        }

    def _call_responses_text(self, instructions, prompt, model=None, max_output_tokens=2048,
                              temperature=0.1, tools=None, timeout=60):
        api_key = self._get_api_key()
        if not api_key:
            raise UserError("API key not found. Set prema_ai.api_key in Settings.")

        model = model or self._resolve_lead_gen_model()
        payload = {
            "model": model,
            "instructions": instructions,
            "input": [{"role": "user", "content": prompt}],
            "max_output_tokens": max_output_tokens,
            "store": False,
        }
        if tools:
            payload["tools"] = tools
        if self._supports_reasoning(model):
            payload["reasoning"] = {"effort": "low"}
        elif temperature is not None:
            payload["temperature"] = temperature

        resp = requests.post(
            "https://api.openai.com/v1/responses",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json=payload,
            timeout=timeout,
        )
        if resp.status_code != 200:
            raise UserError(f"OpenAI error ({resp.status_code}): {resp.text[:200]}")
        return self._extract_response_text(resp.json())

    def _extract_response_text(self, data):
        raw = (data or {}).get("output_text", "") or ""
        if raw:
            return raw.strip()
        for item in (data or {}).get("output", []):
            if item.get("type") != "message":
                continue
            for block in item.get("content", []):
                if block.get("type") == "output_text" and block.get("text"):
                    return block["text"].strip()
        return ""

    def _parse_json_payload(self, raw, expected="object"):
        text = re.sub(r"^```(?:json)?\s*", "", (raw or "").strip())
        text = re.sub(r"\s*```$", "", text.strip())
        if expected == "array":
            start, end = text.find("["), text.rfind("]")
        else:
            start, end = text.find("{"), text.rfind("}")
        if start >= 0 and end >= start:
            text = text[start:end + 1]
        return json.loads(text)

    def _prepare_google_lead_candidates(self, plan):
        center = self._google_geocode_location(plan["location"])
        if not center:
            raise UserError(f"Could not geocode the target area: {plan['location']}")

        seen_ids = {}
        max_candidates = max(plan["count"] * 3, 18)
        for term in plan.get("search_terms", [])[:6]:
            for place in self._google_places_text_search(term, center, plan["radius_km"], plan["location"]):
                place_id = place.get("id")
                if not place_id or place_id in seen_ids:
                    continue
                lat = (place.get("location") or {}).get("latitude")
                lng = (place.get("location") or {}).get("longitude")
                if lat and lng:
                    dist = self._distance_km(center["lat"], center["lng"], lat, lng)
                    if dist > (plan["radius_km"] + 5):
                        continue
                seen_ids[place_id] = {**place, "_search_term": term}
                if len(seen_ids) >= max_candidates:
                    break
            if len(seen_ids) >= max_candidates:
                break

        candidates = []
        for place in list(seen_ids.values())[:max(plan["count"] * 2, 20)]:
            details = self._google_place_details(place["id"]) or place
            candidate = self._candidate_from_place(details, plan, place.get("_search_term", ""))
            if candidate:
                candidates.append(candidate)
        return self._dedupe_lead_candidates(candidates)

    def _google_geocode_location(self, location_text):
        key = self._get_google_maps_api_key()
        resp = requests.get(
            "https://maps.googleapis.com/maps/api/geocode/json",
            params={"address": location_text, "key": key},
            timeout=20,
        )
        resp.raise_for_status()
        data = resp.json()
        result = (data.get("results") or [None])[0]
        if not result:
            return None
        loc = (result.get("geometry") or {}).get("location") or {}
        if loc.get("lat") is None or loc.get("lng") is None:
            return None
        return {
            "lat": float(loc["lat"]),
            "lng": float(loc["lng"]),
            "formatted_address": result.get("formatted_address", location_text),
        }

    def _google_places_text_search(self, term, center, radius_km, location_text):
        north, south, east, west = self._bounding_box(center["lat"], center["lng"], radius_km)
        payload = {
            "textQuery": f"{term} in {location_text}",
            "pageSize": 20,
            "languageCode": "en",
            "locationRestriction": {
                "rectangle": {
                    "low": {"latitude": south, "longitude": west},
                    "high": {"latitude": north, "longitude": east},
                }
            },
        }
        resp = requests.post(
            "https://places.googleapis.com/v1/places:searchText",
            headers={
                "Content-Type": "application/json",
                "X-Goog-Api-Key": self._get_google_maps_api_key(),
                "X-Goog-FieldMask": GOOGLE_TEXT_SEARCH_FIELDS,
            },
            json=payload,
            timeout=30,
        )
        resp.raise_for_status()
        return resp.json().get("places") or []

    def _google_place_details(self, place_id):
        resp = requests.get(
            f"https://places.googleapis.com/v1/places/{place_id}",
            headers={
                "X-Goog-Api-Key": self._get_google_maps_api_key(),
                "X-Goog-FieldMask": GOOGLE_PLACE_DETAIL_FIELDS,
            },
            timeout=30,
        )
        resp.raise_for_status()
        return resp.json()

    def _candidate_from_place(self, place, plan, search_term):
        name = self._place_display_name(place)
        if not name:
            return None

        postal = place.get("postalAddress") or {}
        address_lines = postal.get("addressLines") or []
        city = (
            postal.get("locality")
            or postal.get("sublocality")
            or postal.get("administrativeArea")
            or ""
        )
        state = postal.get("administrativeArea") or ""
        country_code = (postal.get("regionCode") or "").upper()
        country_name = self._country_name_from_code(country_code)
        website = (place.get("websiteUri") or "").strip().rstrip("/")
        phone = (place.get("internationalPhoneNumber") or place.get("nationalPhoneNumber") or "").strip()
        domain = self._normalize_domain(website)
        best_contact = {}

        return {
            "pending_key": place.get("id") or place.get("name") or name,
            "company_name": name,
            "contact_name": best_contact.get("contact_name", ""),
            "email": best_contact.get("email", ""),
            "title": best_contact.get("title", ""),
            "phone": phone or best_contact.get("phone", ""),
            "website": website,
            "domain": domain,
            "street": ", ".join(address_lines) or place.get("formattedAddress") or "",
            "city": city,
            "state": state,
            "zip": postal.get("postalCode") or "",
            "country": country_name or country_code,
            "country_code": country_code,
            "industry": ", ".join((place.get("types") or [])[:3]),
            "description": best_contact.get("reason") or f"Google Places match for '{search_term}'",
            "google_place_id": place.get("id"),
        }

    def _rank_lead_candidates(self, user_msg, candidates, plan):
        ordered = self._dedupe_lead_candidates(
            sorted(candidates, key=self._heuristic_lead_score, reverse=True)
        )
        api_key = self._get_api_key()
        if not api_key or not ordered:
            return ordered

        shortlist = ordered[: max(plan["count"] * 2, 16)]
        payload = [
            {
                "pending_key": lead.get("pending_key"),
                "company_name": lead.get("company_name"),
                "contact_name": lead.get("contact_name"),
                "title": lead.get("title"),
                "email": lead.get("email"),
                "website": lead.get("website"),
                "city": lead.get("city"),
                "state": lead.get("state"),
                "description": lead.get("description"),
                "industry": lead.get("industry"),
            }
            for lead in shortlist
        ]
        prompt = (
            f"User request: {user_msg}\n"
            f"Need the best {plan['count']} leads.\n\n"
            f"Candidate companies:\n{json.dumps(payload, indent=2)}\n\n"
            "Return ONLY a JSON array ordered best-to-worst. "
            "Each item must contain: pending_key, score, reason. "
            "Prefer shippers, distributors, warehouses, food/produce suppliers, and companies likely to tender freight. "
            "Avoid obvious retail storefront noise when better logistics targets exist."
        )
        try:
            raw = self._call_responses_text(
                instructions="Rank B2B freight lead candidates and return only JSON.",
                prompt=prompt,
                model=self._resolve_lead_gen_model(),
                max_output_tokens=1800,
                temperature=0.1,
                timeout=45,
            )
            ranked = self._parse_json_payload(raw, expected="array") or []
            if not isinstance(ranked, list):
                return ordered
            by_key = {
                str(lead.get("pending_key")): lead
                for lead in shortlist
                if lead.get("pending_key")
            }
            selected = []
            selected_keys = set()
            for row in ranked:
                pending_key = str(row.get("pending_key") or "")
                if not pending_key or pending_key in selected_keys:
                    continue
                lead = by_key.get(pending_key)
                if not lead:
                    continue
                updated = dict(lead)
                if row.get("reason"):
                    updated["description"] = str(row["reason"]).strip()
                selected.append(updated)
                selected_keys.add(pending_key)
            if selected:
                return selected + [
                    lead for lead in ordered
                    if str(lead.get("pending_key") or "") not in selected_keys
                ]
        except Exception as exc:
            _logger.warning("Lead ranking fallback: %s", exc)
        return ordered

    def _lead_candidate_dedupe_key(self, lead):
        place_id = str(lead.get("google_place_id") or lead.get("pending_key") or "").strip()
        if place_id:
            return f"place:{place_id}"

        domain = self._normalize_domain(
            lead.get("domain")
            or lead.get("website")
            or self._extract_domain_from_email(lead.get("email"))
        )
        if domain:
            return f"domain:{domain}"

        parts = [
            (lead.get("company_name") or "").strip().lower(),
            (lead.get("street") or "").strip().lower(),
            (lead.get("city") or "").strip().lower(),
            (lead.get("state") or "").strip().lower(),
        ]
        return "name:" + "|".join(re.sub(r"\s+", " ", part) for part in parts)

    def _dedupe_lead_candidates(self, candidates):
        deduped = []
        seen = set()
        for idx, lead in enumerate(candidates or []):
            normalized = dict(lead or {})
            if not normalized.get("pending_key"):
                normalized["pending_key"] = (
                    normalized.get("google_place_id")
                    or normalized.get("name")
                    or f"candidate-{idx + 1}"
                )
            key = self._lead_candidate_dedupe_key(normalized)
            if key in seen:
                continue
            seen.add(key)
            deduped.append(normalized)
        return deduped

    def _heuristic_lead_score(self, lead):
        title = (lead.get("title") or "").lower()
        description = (lead.get("description") or "").lower()
        score = 0
        if lead.get("email"):
            score += 35
        if lead.get("contact_name"):
            score += 12
        if lead.get("website"):
            score += 10
        if any(word in title for word in ["director", "vp", "vice president", "head"]):
            score += 22
        elif "manager" in title:
            score += 18
        elif any(word in title for word in ["coordinator", "supervisor", "dispatch"]):
            score += 12
        if any(word in description for word in ["warehouse", "distribution", "logistics", "supply", "produce", "food"]):
            score += 10
        return score

    def _get_lead_gen_titles(self):
        titles = []
        try:
            profile = self.env["premafirm.business.profile"].sudo().get_profile()
            titles = [title for title in profile.icp_title_ids.mapped("name") if title]
        except Exception:
            pass
        titles.extend(LEAD_GEN_DEFAULT_TITLES)
        deduped = []
        seen = set()
        for title in titles:
            key = title.strip().lower()
            if key and key not in seen:
                seen.add(key)
                deduped.append(title.strip())
        return deduped

    def _normalize_domain(self, value):
        if not value:
            return ""
        text = str(value).strip().lower()
        text = re.sub(r"^https?://", "", text)
        text = re.sub(r"^www\.", "", text)
        text = text.split("/")[0].split("?")[0]
        return text.strip()

    def _extract_domain_from_email(self, email):
        if email and "@" in email:
            return email.split("@", 1)[1].strip().lower()
        return ""

    def _place_display_name(self, place):
        display = (place or {}).get("displayName")
        if isinstance(display, dict):
            return (display.get("text") or "").strip()
        return str(display or "").strip()

    def _country_name_from_code(self, code):
        if not code:
            return ""
        country = self.env["res.country"].sudo().search([("code", "=", code.upper())], limit=1)
        return country.name if country else code.upper()

    def _bounding_box(self, lat, lng, radius_km):
        lat_delta = radius_km / 111.32
        lng_delta = radius_km / max(111.32 * math.cos(math.radians(lat)), 1e-6)
        return lat + lat_delta, lat - lat_delta, lng + lng_delta, lng - lng_delta

    def _distance_km(self, lat1, lng1, lat2, lng2):
        r = 6371.0
        dlat = math.radians(lat2 - lat1)
        dlng = math.radians(lng2 - lng1)
        a = (
            math.sin(dlat / 2) ** 2
            + math.cos(math.radians(lat1))
            * math.cos(math.radians(lat2))
            * math.sin(dlng / 2) ** 2
        )
        return r * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))

    # ── Tag creation via chat ──────────────────────────────────────────
    def _handle_tag_creation(self, user_msg):
        """Create CRM or partner tags from a chat instruction."""
        import re as _re
        # Extract tag name(s) from message
        m = _re.search(
            r'(?:create|add|make|new)\s+tag\s+["\']?([^"\']+)["\']?',
            user_msg, _re.IGNORECASE
        )
        if not m:
            return "Please specify the tag name, e.g. \"Create tag Warm Lead\""

        tag_name = m.group(1).strip().strip('"\'')

        # Create in crm.tag (for CRM leads) and res.partner.category (for contacts)
        created = []
        for model, label in [("crm.tag", "CRM"), ("res.partner.category", "Contact")]:
            existing = self.env[model].sudo().search([("name", "=ilike", tag_name)], limit=1)
            if not existing:
                self.env[model].sudo().create({"name": tag_name})
                created.append(label)

        if created:
            return f"Tag **{tag_name}** created in: {', '.join(created)}. You can now apply it to CRM leads and contacts."
        return f"Tag **{tag_name}** already exists."

    # ── Lead creation confirmation ───────────────────────────────────
    @api.model
    def confirm_create_leads(self, session_id):
        """Called when user clicks 'Create Leads'. Creates records from pending data."""
        session = self.browse(session_id)
        if not session.exists():
            return {"error": "Session not found"}
        if not session.pending_lead_data:
            return {"error": "No pending leads to create"}

        leads = session._load_pending_leads()
        created_ids, updated_ids, existing_ids, errors = [], [], [], []
        for ld in leads:
            try:
                result = session._upsert_lead_from_ai(ld)
                lead = result["lead"]
                if result["status"] == "created":
                    created_ids.append(lead.id)
                elif result["status"] == "updated":
                    updated_ids.append(lead.id)
                else:
                    existing_ids.append(lead.id)
            except Exception as exc:
                _logger.warning("Lead import failed for %s: %s", ld.get("company_name"), exc)
                errors.append(f"{ld.get('company_name', '?')}: {exc}")

        session.pending_lead_data = False
        msg = (
            f"Lead import complete.\n"
            f"Created: {len(created_ids)}\n"
            f"Updated existing leads: {len(updated_ids)}\n"
            f"Already in CRM: {len(existing_ids)}"
        )
        if errors:
            msg += f"\nErrors: {len(errors)} — {'; '.join(errors[:3])}"
        self.env["prema.ai.message"].sudo().create({
            "session_id": session_id, "role": "assistant", "content": msg})
        return {
            "created": len(created_ids),
            "updated": len(updated_ids),
            "existing": len(existing_ids),
            "errors": errors,
            "message": msg,
            "pending_action": None,
        }

    @api.model
    def confirm_create_single_lead(self, session_id, pending_key):
        session = self.browse(session_id)
        if not session.exists():
            return {"error": "Session not found"}

        leads = session._load_pending_leads()
        target_idx = next(
            (idx for idx, lead in enumerate(leads) if str(lead.get("pending_key")) == str(pending_key)),
            None,
        )
        if target_idx is None:
            return {"error": "Lead card not found"}

        lead_data = leads[target_idx]
        try:
            result = session._upsert_lead_from_ai(lead_data)
        except Exception as exc:
            return {"error": str(exc), "pending_action": session._build_pending_action_payload()}

        remaining = leads[:target_idx] + leads[target_idx + 1:]
        session._save_pending_leads(remaining)

        if result["status"] == "created":
            msg = f"Created CRM lead for {lead_data.get('company_name')}."
        elif result["status"] == "updated":
            msg = f"Updated the existing CRM lead for {lead_data.get('company_name')} with the new contact data."
        else:
            msg = f"{lead_data.get('company_name')} is already in CRM. No duplicate lead was created."

        self.env["prema.ai.message"].sudo().create({
            "session_id": session_id, "role": "assistant", "content": msg})
        return {
            "created": 1 if result["status"] == "created" else 0,
            "updated": 1 if result["status"] == "updated" else 0,
            "existing": 1 if result["status"] == "existing" else 0,
            "message": msg,
            "pending_action": session._build_pending_action_payload(),
        }

    @api.model
    def cancel_pending_action(self, session_id):
        """Cancel any pending confirmation action on the session."""
        session = self.browse(session_id)
        if not session.exists():
            return {"error": "Session not found"}
        session.pending_lead_data = False
        msg = "Lead creation cancelled. No records were created."
        self.env["prema.ai.message"].sudo().create({
            "session_id": session_id, "role": "assistant", "content": msg})
        return {"cancelled": True, "message": msg, "pending_action": None}

    def create_lead_from_ai(self, data):
        return self._upsert_lead_from_ai(data)["lead"]

    def _preview_crm_target(self, data):
        match = self._match_existing_crm_records(data)
        if match["lead"] and match["lead_has_contact"]:
            return {
                "existing_company_id": match["company_partner"].id if match["company_partner"] else False,
                "existing_lead_id": match["lead"].id,
                "skip_create": True,
                "import_label": "Already in CRM",
                "status_note": f"Already on CRM lead #{match['lead'].id} {match['lead'].name}",
            }
        if match["lead"]:
            return {
                "existing_company_id": match["company_partner"].id if match["company_partner"] else False,
                "existing_lead_id": match["lead"].id,
                "skip_create": False,
                "import_label": "Update Existing Lead",
                "status_note": f"Will update CRM lead #{match['lead'].id} {match['lead'].name}",
            }
        if match["company_partner"]:
            return {
                "existing_company_id": match["company_partner"].id,
                "existing_lead_id": False,
                "skip_create": False,
                "import_label": "Add to CRM",
                "status_note": f"Will reuse company {match['company_partner'].name}",
            }
        return {
            "existing_company_id": False,
            "existing_lead_id": False,
            "skip_create": False,
            "import_label": "Add to CRM",
            "status_note": "New company and lead",
        }

    def _upsert_lead_from_ai(self, data):
        self.ensure_one()
        company_name = (data.get("company_name") or "").strip()
        if not company_name or company_name.lower() == "unknown":
            raise UserError("Company name is required to create a lead.")

        match = self._match_existing_crm_records(data)
        company = match["company_partner"]
        contact = match["contact_partner"]
        lead = match["lead"]
        country = self._resolve_country(data.get("country_code") or data.get("country"))
        state = self._resolve_state(data.get("state_code") or data.get("state"), country)

        company_vals = self._build_company_partner_vals(data, country, state)
        if company:
            self._update_partner_missing_fields(company, company_vals)
        else:
            company = self.env["res.partner"].sudo().create(company_vals)

        if not contact:
            contact_vals = self._build_contact_partner_vals(data, company, country, state)
            if contact_vals:
                contact = self.env["res.partner"].sudo().create(contact_vals)
        elif company and contact.parent_id != company:
            contact.parent_id = company.id
            self._update_partner_missing_fields(contact, self._build_contact_partner_vals(data, company, country, state))
        else:
            self._update_partner_missing_fields(contact, self._build_contact_partner_vals(data, company, country, state))

        if lead and match["lead_has_contact"]:
            self._update_existing_lead_fields(lead, data, company, contact, country, state)
            return {"status": "existing", "lead": lead, "company": company, "contact": contact}

        if lead:
            self._update_existing_lead_fields(lead, data, company, contact, country, state)
            return {"status": "updated", "lead": lead, "company": company, "contact": contact}

        lead_vals = {
            "name": company.name,
            "type": "opportunity",
            "stage_id": self._new_stage_id(),
            "user_id": self.user_id.id or self.env.user.id,
            "partner_id": contact.id if contact else company.id,
            "partner_name": company.name,
            "contact_name": contact.name if contact else (data.get("contact_name") or ""),
            "email_from": (data.get("email") or (contact.email if contact else "") or "").strip(),
            "phone": (data.get("phone") or (contact.phone if contact else "") or company.phone or "").strip(),
            "website": company.website or data.get("website") or "",
            "street": company.street or data.get("street") or "",
            "city": company.city or data.get("city") or "",
            "zip": company.zip or data.get("zip") or "",
            "description": data.get("description") or "",
        }
        if country:
            lead_vals["country_id"] = country.id
        if state:
            lead_vals["state_id"] = state.id
        if "outreach_stage" in self.env["crm.lead"]._fields:
            lead_vals["outreach_stage"] = "new"

        lead = self.env["crm.lead"].sudo().create(lead_vals)
        self._apply_lead_tags(lead, data)
        return {"status": "created", "lead": lead, "company": company, "contact": contact}

    def _match_existing_crm_records(self, data):
        company = self._find_company_partner(data)
        contact = self._find_contact_partner(company, data)
        if not company and contact and contact.parent_id:
            company = contact.parent_id
        lead = self._find_open_company_lead(company, data)
        lead_has_contact = False

        email = (data.get("email") or "").strip().lower()
        contact_name = (data.get("contact_name") or "").strip().lower()
        title = (data.get("title") or data.get("job_title") or "").strip().lower()
        if lead:
            if email and (lead.email_from or "").strip().lower() == email:
                lead_has_contact = True
            if not lead_has_contact and contact and lead.partner_id == contact:
                lead_has_contact = True
        return {
            "company_partner": company,
            "contact_partner": contact,
            "lead": lead,
            "lead_has_contact": lead_has_contact,
        }

    def _find_company_partner(self, data):
        Partner = self.env["res.partner"].sudo()
        company_name = (data.get("company_name") or "").strip()
        city = (data.get("city") or "").strip().lower()
        domain = self._normalize_domain(
            data.get("domain") or data.get("website") or self._extract_domain_from_email(data.get("email"))
        )

        if domain:
            candidates = Partner.search([("active", "=", True), ("website", "ilike", domain)])
            company = candidates.filtered(
                lambda partner: partner.is_company and self._normalize_domain(partner.website) == domain
            )[:1]
            if company:
                return company

            email_contacts = Partner.search([("active", "=", True), ("email", "ilike", f"@{domain}")], limit=10)
            for contact in email_contacts:
                if contact.parent_id:
                    return contact.parent_id
                if contact.is_company:
                    return contact

        if not company_name:
            return False
        companies = Partner.search([
            ("active", "=", True),
            ("is_company", "=", True),
            ("name", "=ilike", company_name),
        ])
        if city:
            city_matches = companies.filtered(lambda company: (company.city or "").strip().lower() == city)
            if city_matches:
                return city_matches[0]
        return companies[:1]

    def _find_contact_partner(self, company, data):
        Partner = self.env["res.partner"].sudo()
        email = (data.get("email") or "").strip()
        if email:
            contact = Partner.search([("active", "=", True), ("email", "=ilike", email)], limit=1)
            if contact and not contact.is_company and (not company or contact.parent_id == company):
                return contact

        contact_name = (data.get("contact_name") or "").strip()
        if contact_name and company:
            candidates = Partner.search([
                ("active", "=", True),
                ("parent_id", "=", company.id),
                ("name", "=ilike", contact_name),
            ], limit=5)
            title = (data.get("title") or data.get("job_title") or "").strip().lower()
            if title:
                match = candidates.filtered(lambda partner: (partner.function or "").strip().lower() == title)[:1]
                if match:
                    return match
            if candidates:
                return candidates[0]
        return False

    def _find_open_company_lead(self, company, data):
        Lead = self.env["crm.lead"].sudo()
        email = (data.get("email") or "").strip()
        if email:
            lead = Lead.search([("active", "=", True), ("email_from", "=ilike", email)], order="create_date desc", limit=1)
            if lead:
                return lead

        domain = self._normalize_domain(
            data.get("domain") or data.get("website") or self._extract_domain_from_email(email)
        )
        company_name = (data.get("company_name") or "").strip()
        if company:
            lead = Lead.search([
                ("active", "=", True),
                "|",
                "|",
                ("partner_id", "=", company.id),
                ("partner_id.parent_id", "=", company.id),
                ("partner_name", "=ilike", company.name),
            ], order="create_date desc", limit=1)
            if lead:
                return lead

        if domain:
            candidates = Lead.search([("active", "=", True), ("website", "ilike", domain)], order="create_date desc", limit=5)
            filtered = candidates.filtered(lambda lead: self._normalize_domain(lead.website) == domain)
            if filtered:
                return filtered[0]

        if company_name:
            return Lead.search([("active", "=", True), ("partner_name", "=ilike", company_name)], order="create_date desc", limit=1)
        return False

    def _resolve_country(self, value):
        if not value:
            return False
        Country = self.env["res.country"].sudo()
        country = Country.search([("code", "=", str(value).strip().upper())], limit=1)
        if country:
            return country
        return Country.search([("name", "=ilike", str(value).strip())], limit=1)

    def _resolve_state(self, value, country=False):
        if not value:
            return False
        State = self.env["res.country.state"].sudo()
        domain = [("country_id", "=", country.id)] if country else []
        state = State.search(domain + [("code", "=ilike", str(value).strip())], limit=1)
        if state:
            return state
        return State.search(domain + [("name", "=ilike", str(value).strip())], limit=1)

    def _build_company_partner_vals(self, data, country=False, state=False):
        vals = {
            "name": (data.get("company_name") or "").strip(),
            "is_company": True,
            "website": (data.get("website") or "").strip(),
            "phone": (data.get("phone") or "").strip(),
            "street": (data.get("street") or "").strip(),
            "city": (data.get("city") or "").strip(),
            "zip": (data.get("zip") or "").strip(),
        }
        if country:
            vals["country_id"] = country.id
        if state:
            vals["state_id"] = state.id
        return vals

    def _build_contact_partner_vals(self, data, company, country=False, state=False):
        contact_name = (data.get("contact_name") or "").strip()
        email = (data.get("email") or "").strip()
        title = (data.get("title") or data.get("job_title") or "").strip()
        phone = (data.get("phone") or "").strip()
        if not any([contact_name, email, title]):
            return {}
        vals = {
            "name": contact_name or email or f"{company.name} Contact",
            "parent_id": company.id if company else False,
            "type": "contact",
            "email": email,
            "phone": phone,
            "function": title,
            "city": (data.get("city") or "").strip(),
        }
        if country:
            vals["country_id"] = country.id
        if state:
            vals["state_id"] = state.id
        return vals

    def _update_partner_missing_fields(self, partner, vals):
        updates = {}
        for field_name, value in (vals or {}).items():
            if value and not partner[field_name]:
                updates[field_name] = value
        if updates:
            partner.write(updates)
        return updates

    def _update_existing_lead_fields(self, lead, data, company, contact, country=False, state=False):
        vals = {}
        if contact and not lead.partner_id:
            vals["partner_id"] = contact.id
        if company and not lead.partner_name:
            vals["partner_name"] = company.name
        if contact and not lead.contact_name:
            vals["contact_name"] = contact.name
        if data.get("email") and not lead.email_from:
            vals["email_from"] = data.get("email")
        if company.website and not lead.website:
            vals["website"] = company.website
        if data.get("phone") and not lead.phone:
            vals["phone"] = data.get("phone")
        if company.street and not lead.street:
            vals["street"] = company.street
        if company.city and not lead.city:
            vals["city"] = company.city
        if company.zip and not lead.zip:
            vals["zip"] = company.zip
        if country and not lead.country_id:
            vals["country_id"] = country.id
        if state and not lead.state_id:
            vals["state_id"] = state.id
        if data.get("description") and not lead.description:
            vals["description"] = data.get("description")
        if self.user_id and not lead.user_id:
            vals["user_id"] = self.user_id.id
        if vals:
            lead.write(vals)

    def _apply_lead_tags(self, lead, data):
        tag_names = data.get("tags") or []
        if isinstance(tag_names, str):
            tag_names = [tag.strip() for tag in tag_names.split(",") if tag.strip()]
        if not tag_names:
            return
        Tag = self.env.get("crm.tag")
        if Tag is None:
            return
        for tag_name in tag_names:
            tag = Tag.sudo().search([("name", "=ilike", tag_name)], limit=1)
            if not tag:
                tag = Tag.sudo().create({"name": tag_name})
            lead.sudo().write({"tag_ids": [(4, tag.id)]})

    def _new_stage_id(self):
        stage = self.env["crm.stage"].sudo().search([("name", "=", "New")], limit=1)
        return stage.id if stage else False

    # ── API key + model resolution ──────────────────────────────────
    def _get_api_key(self):
        p = self.env["ir.config_parameter"].sudo()
        return (p.get_param("prema_ai.api_key") or p.get_param("openai.api_key") or "").strip()

    def _resolve_model(self, mode="primary"):
        p = self.env["ir.config_parameter"].sudo()
        if mode == "fast":
            return p.get_param("prema_ai.fast_model", "gpt-4.1")
        elif mode == "vision":
            return p.get_param("prema_ai.vision_model", "gpt-4o")
        return p.get_param("prema_ai.primary_model", "gpt-4o-mini")

    @staticmethod
    def _supports_temperature(model):
        """Reasoning models (o-series) and gpt-5.5 reject the temperature parameter."""
        import re
        m = (model or "").lower().strip()
        return not (re.match(r"^o\d", m) or m in {"gpt-5.5", "gpt-5"})

    @staticmethod
    def _supports_reasoning(model):
        import re
        m = (model or "").lower().strip()
        return bool(re.match(r"^o\d", m) or m.startswith("gpt-5"))

    # ====================================================================
    # File content extraction — handles every attachment type
    # ====================================================================
    def _extract_b64_file_content(self, b64, mime, filename):
        """Extract readable text from any base64-encoded file.
        Handles: PDF (text layer + OCR), Excel, CSV, Word, plain text.
        Returns extracted string (up to 8000 chars) or empty string on failure.
        """
        try:
            import io as _io
            file_bytes = base64.b64decode(b64)
            mime = (mime or "").lower()
            name = (filename or "").lower()

            # ── PDF ──────────────────────────────────────────────────
            if "pdf" in mime or name.endswith(".pdf"):
                # 1) Text layer (fast, works for digital PDFs)
                try:
                    import pdfplumber
                    parts = []
                    with pdfplumber.open(_io.BytesIO(file_bytes)) as pdf:
                        for page in pdf.pages[:15]:
                            t = page.extract_text(layout=True) or ""
                            if t.strip():
                                parts.append(t)
                    text = "\n".join(parts).strip()
                    if text:
                        return text[:8000]
                except Exception:
                    pass
                # 2) OCR fallback for scanned PDFs
                try:
                    from pdf2image import convert_from_bytes
                    import pytesseract
                    imgs = convert_from_bytes(file_bytes, dpi=150, first_page=1, last_page=6)
                    texts = [pytesseract.image_to_string(img) for img in imgs]
                    text = "\n".join(t for t in texts if t).strip()
                    if text:
                        return text[:8000]
                except Exception:
                    pass
                return ""

            # ── Excel ─────────────────────────────────────────────────
            if ("spreadsheet" in mime or "excel" in mime
                    or name.endswith(".xlsx") or name.endswith(".xls")):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(
                        _io.BytesIO(file_bytes), read_only=True, data_only=True)
                    lines = []
                    for sheet in list(wb.worksheets)[:5]:
                        lines.append(f"=== Sheet: {sheet.title} ===")
                        for row in sheet.iter_rows(max_row=300, values_only=True):
                            line = "\t".join(
                                str(c) if c is not None else "" for c in row)
                            if line.strip():
                                lines.append(line)
                    return "\n".join(lines)[:8000]
                except Exception:
                    pass

            # ── CSV ───────────────────────────────────────────────────
            if "csv" in mime or name.endswith(".csv"):
                try:
                    return file_bytes.decode("utf-8", errors="replace")[:8000]
                except Exception:
                    pass

            # ── Word document ─────────────────────────────────────────
            if ("word" in mime or "officedocument.wordprocessing" in mime
                    or name.endswith(".docx") or name.endswith(".doc")):
                try:
                    from docx import Document as DocxDoc
                    doc = DocxDoc(_io.BytesIO(file_bytes))
                    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                    # Also include tables
                    for tbl in doc.tables:
                        for row in tbl.rows:
                            row_text = "\t".join(
                                cell.text.strip() for cell in row.cells)
                            if row_text.strip():
                                text += "\n" + row_text
                    return text[:8000]
                except Exception:
                    pass

            # ── Plain text / markdown / log ───────────────────────────
            if "text/" in mime or name.endswith((".txt", ".md", ".log", ".json")):
                try:
                    return file_bytes.decode("utf-8", errors="replace")[:8000]
                except Exception:
                    pass

            return ""
        except Exception as ex:
            _logger.debug("File content extraction failed (%s): %s", filename, ex)
            return ""

    # ====================================================================
    # ERP CONTEXT BUILDER — full database read across ALL modules
    # ====================================================================
    def _gather_erp_context(self, user_msg, chat_mode="standard",
                            has_attachments=False, skip_att_scan=False):
        lower = (user_msg or "").lower()
        sections = []

        try:
            sections.append(self._ctx_system_summary())

            entity = self._extract_entity_name(user_msg)
            keywords = self._extract_keywords(lower)

            if entity:
                sections.append(self._ctx_partner_detail(entity))

            # Financial context — only when relevant keywords present
            _finance_words = [
                "invoice", "bill", "payment", "paid", "owing", "outstanding",
                "receivable", "payable", "revenue", "expense", "account",
                "financial", "finance", "money", "charge", "amount", "balance",
                "credit", "debit", "vendor", "supplier", "customer invoice",
            ]
            if any(w in lower for w in _finance_words) or entity:
                sections.append(self._ctx_vendor_bills(entity))
                sections.append(self._ctx_customer_invoices(entity))

            # Keyword search in bill lines
            if keywords:
                sections.append(self._ctx_search_bill_lines(keywords))

            # Module-specific context
            if any(w in lower for w in ["subscription", "recurring", "insurance"]):
                sections.append(self._ctx_subscriptions(entity))
            if any(w in lower for w in ["lead", "crm", "pipeline", "prospect"]):
                sections.append(self._ctx_crm())
            if any(w in lower for w in ["task", "project"]):
                sections.append(self._ctx_tasks())
            if any(w in lower for w in ["expense"]):
                sections.append(self._ctx_expenses())
            if any(w in lower for w in ["sale", "order", "quotation"]):
                sections.append(self._ctx_sales())
            if any(w in lower for w in ["journal", "entry", "debit", "credit"]):
                sections.append(self._ctx_journal_entries(entity))
            if any(w in lower for w in ["payment", "paid"]):
                sections.append(self._ctx_payments(entity))
            if any(w in lower for w in ["document", "attachment", "file", "pod"]):
                sections.append(self._ctx_documents(entity))

            # Attachment content reading — triggered by document/image/file keywords across ALL models
            _doc_words = [
                "shipper", "shippers", "consignee", "bill of lading", "bol",
                "invoice attachment", "attached file", "attached document",
                "read attachment", "scan attachment", "analyze attachment",
                "review attachment", "go through attachment",
                "from attachment", "in the attachment",
                "check attachment", "see attachment", "find in attachment",
                "open attachment", "view attachment", "look at attachment",
                "attachment for", "documents for customer", "files for",
                "who are they shipping", "who is shipping", "who is the shipper",
                "best customer", "top customer", "biggest customer",
                "who did", "who hauled", "what shippers",
                "photo", "picture", "image", "scan", "scanned",
                "contract", "agreement", "certificate", "insurance doc",
                "pod", "proof of delivery", "signed", "signature",
            ]
            _read_att = (
                any(w in lower for w in _doc_words)
                or ("go through" in lower and any(
                    w in lower for w in ["invoice", "attachment", "file", "document", "bill"]))
                or ("haul" in lower and any(
                    w in lower for w in ["list", "who", "shipper", "for", "them", "customer"]))
                or ("list" in lower and any(
                    w in lower for w in ["shipper", "consignee", "carrier", "hauled", "haul"]))
                or ("check" in lower and any(
                    w in lower for w in ["attachment", "attached", "file", "document", "photo", "picture"]))
                or ("for customer" in lower and any(
                    w in lower for w in ["attachment", "ship", "shipper", "haul"]))
                or ("analyze" in lower or "analyse" in lower)
                or ("read" in lower and any(
                    w in lower for w in ["file", "document", "attachment", "invoice", "contract"]))
                or chat_mode == "document_qa"
            )
            if _read_att and not skip_att_scan:
                # Skip stored-attachment OCR when user uploaded a file to analyze directly
                att_entity = entity or self._extract_entity_name_fuzzy(user_msg)
                sections.append(self._ctx_invoice_attachments(att_entity))
                sections.append(self._ctx_all_model_attachments(att_entity))

            # Fleet module
            if any(w in lower for w in [
                "fleet", "truck", "vehicle", "driver", "fuel", "diesel",
                "maintenance", "mileage", "km", "odometer",
            ]):
                sections.append(self._ctx_fleet())

            # Purchase orders and lines
            if any(w in lower for w in ["purchase", "po", "purchase order"]):
                sections.append(self._ctx_purchases(entity))
                sections.append(self._ctx_purchase_lines(entity))

            # Employees / HR
            if any(w in lower for w in ["employee", "staff", "worker", "hr", "human resource"]):
                sections.append(self._ctx_employees())

            # Stock / delivery / shipments
            if any(w in lower for w in [
                "delivery", "shipment", "pickup", "receipt", "transfer",
                "warehouse", "stock", "inventory", "dispatch",
            ]):
                sections.append(self._ctx_stock_pickings(entity))

            # Account balances
            if any(w in lower for w in [
                "balance", "account balance", "receivable", "payable",
                "cash", "bank", "credit card",
            ]):
                sections.append(self._ctx_account_balances())

            # Deep audit
            if any(w in lower for w in [
                "summary", "detailed", "check", "audit", "review",
                "wrong", "fix", "analyze", "issue", "report",
            ]):
                if entity:
                    sections.append(self._ctx_partner_full_audit(entity))
                sections.append(self._ctx_accounting_health())

            # WA Negotiations
            if any(w in lower for w in [
                "negotiat", "dispatcher", "broker", "their offer", "our counter",
                "agreed rate", "rate confirmation", "load tender", "wa deal",
                "wa negotiation", "whatsapp deal", "thomas", "dispatch",
            ]):
                sections.append(self._ctx_wa_negotiations(entity))

            # Pricing Review: inject pricing engine history + rules
            if chat_mode == "pricing_review":
                sections.append(self._ctx_pricing_context())

            # Source reference files from Chat Settings
            src = self._ctx_source_library()
            if src:
                sections.append(src)

        except Exception as e:
            _logger.error("Context error: %s\n%s", e, traceback.format_exc())
            sections.append(f"[Context error: {e}]")

        return "\n\n".join([s for s in sections if s and s.strip()])

    def _extract_entity_name(self, msg):
        patterns = [
            r"(?:for|about|from|with|check|review|regarding)\s+([A-Z][A-Za-z\s&.,'-]+?)(?:\s*$|\s*\?|\s*\.|\s+and\s|\s+if\s)",
            r"([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,4})",
        ]
        stops = {"the","and","for","what","how","can","you","please","need","show",
                 "find","tell","give","get","all","check","review","look","help",
                 "much","have","this","that","year","total","yes","just","create",
                 "make","generate","export","download","research"}
        for pat in patterns:
            m = re.search(pat, msg or "")
            if m:
                name = m.group(1).strip().rstrip(".,;:")
                if len(name) > 2 and name.lower() not in stops:
                    return name
        return None

    def _extract_entity_name_fuzzy(self, msg):
        """Case-insensitive entity extraction for names like 'jamal' or 'smith'.
        Looks for patterns like 'for/named/customer/client <name>'.
        """
        patterns = [
            r"(?:for|named|customer|client|contact|company|about|from|with)\s+([A-Za-z][A-Za-z\s&.,'-]{1,40}?)(?:\s+to|\s+and|\s*$|\s*\?|\s*\.)",
            r"([A-Za-z][a-z]{2,}(?:\s+[A-Za-z][a-z]{2,}){0,3})",
        ]
        stops = {"the","and","for","what","how","can","you","please","need","show","find",
                 "tell","give","get","all","check","review","look","help","much","have",
                 "this","that","year","total","yes","just","create","make","generate",
                 "export","download","research","invoice","attachment","customer","client",
                 "named","best","good","top","list","who","they","their","about","from",
                 "with","same","also","more","less","when","then","than","been","were"}
        for pat in patterns:
            m = re.search(pat, msg or "", re.IGNORECASE)
            if m:
                name = m.group(1).strip().rstrip(".,;:")
                if len(name) > 2 and name.lower() not in stops:
                    return name.title()  # capitalise for DB ilike search
        return None

    def _extract_keywords(self, lower):
        kw_list = [
            "diesel", "fuel", "gas", "petrol", "maintenance", "repair",
            "insurance", "tire", "licence", "license", "permit", "toll",
            "scale", "parking", "wash", "trailer", "truck", "reefer",
            "dispatch", "broker", "freight", "shipping", "delivery",
            "office", "phone", "internet", "rent", "utility",
        ]
        return [w for w in kw_list if w in lower]

    # ── Context builders ────────────────────────────────────────────
    def _ctx_system_summary(self):
        try:
            e = self.env
            return (
                "=== ERP SYSTEM SUMMARY ===\n"
                f"Company: {e.company.name}\n"
                f"Vendor bills: {e['account.move'].sudo().search_count([('move_type','=','in_invoice')])}\n"
                f"Customer invoices: {e['account.move'].sudo().search_count([('move_type','=','out_invoice')])}\n"
                f"Draft bills: {e['account.move'].sudo().search_count([('move_type','=','in_invoice'),('state','=','draft')])}\n"
                f"Active vendors: {e['res.partner'].sudo().search_count([('supplier_rank','>',0)])}\n"
                f"Active customers: {e['res.partner'].sudo().search_count([('customer_rank','>',0)])}\n"
                f"CRM leads: {e['crm.lead'].sudo().search_count([])}")
        except Exception as ex:
            return f"[Summary error: {ex}]"

    def _ctx_partner_detail(self, name):
        if not name:
            return ""
        try:
            ps = self.env["res.partner"].sudo().search_read(
                [("name", "ilike", name), ("active", "=", True)],
                ["id", "name", "email", "phone", "supplier_rank",
                 "customer_rank", "city", "country_id", "vat"],
                limit=5)
            if not ps:
                return f"=== PARTNER '{name}' ===\nNot found."
            lines = [f"=== PARTNERS matching '{name}' ==="]
            for p in ps:
                lines.append(
                    f"  ID {p['id']}: {p['name']} | email: {p.get('email') or '-'} | "
                    f"phone: {p.get('phone') or '-'} | vendor: {p.get('supplier_rank',0)} | "
                    f"customer: {p.get('customer_rank',0)} | city: {p.get('city') or '-'}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Partner error: {ex}]"

    def _ctx_vendor_bills(self, pn=None):
        try:
            dom = [("move_type", "=", "in_invoice")]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            bills = self.env["account.move"].sudo().search_read(
                dom, ["id", "name", "partner_id", "ref", "invoice_date",
                       "invoice_date_due", "amount_total", "amount_residual",
                       "amount_tax", "state", "payment_state"],
                order="invoice_date desc", limit=30)
            lbl = f"for '{pn}'" if pn else "(recent 30)"
            if not bills:
                return f"=== VENDOR BILLS {lbl} ===\nNone."
            lines = [f"=== VENDOR BILLS {lbl} - {len(bills)} records ==="]
            for b in bills:
                lines.append(
                    f"  {b['name']} | {b['partner_id'][1] if b['partner_id'] else '-'} | "
                    f"ref: {b.get('ref') or '-'} | date: {b.get('invoice_date') or '-'} | "
                    f"due: {b.get('invoice_date_due') or '-'} | "
                    f"total: ${b['amount_total']:.2f} | tax: ${b.get('amount_tax',0):.2f} | "
                    f"owing: ${b['amount_residual']:.2f} | {b['state']} | {b.get('payment_state') or '-'}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Bills error: {ex}]"

    def _ctx_customer_invoices(self, pn=None):
        try:
            dom = [("move_type", "=", "out_invoice")]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            invs = self.env["account.move"].sudo().search_read(
                dom, ["id", "name", "partner_id", "invoice_date",
                       "amount_total", "amount_residual", "state", "payment_state"],
                order="invoice_date desc", limit=20)
            lbl = f"for '{pn}'" if pn else "(recent 20)"
            if not invs:
                return f"=== INVOICES {lbl} ===\nNone."
            lines = [f"=== CUSTOMER INVOICES {lbl} - {len(invs)} ==="]
            for i in invs:
                lines.append(
                    f"  {i['name']} | {i['partner_id'][1] if i['partner_id'] else '-'} | "
                    f"date: {i.get('invoice_date') or '-'} | "
                    f"${i['amount_total']:.2f} | owing: ${i['amount_residual']:.2f} | "
                    f"{i['state']} | {i.get('payment_state') or '-'}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Invoices error: {ex}]"

    def _ctx_search_bill_lines(self, keywords):
        try:
            domain = [("move_id.move_type", "in", ["in_invoice", "in_refund"]),
                      ("move_id.state", "=", "posted")]
            kw_dom = ["|"] * (len(keywords) - 1) if len(keywords) > 1 else []
            for kw in keywords:
                kw_dom.append(("name", "ilike", kw))
            lines = self.env["account.move.line"].sudo().search_read(
                domain + kw_dom,
                ["id", "move_id", "name", "product_id", "account_id",
                 "quantity", "price_unit", "debit", "credit", "partner_id", "date"],
                order="date desc", limit=50)
            if not lines:
                return f"=== BILL LINES {keywords} ===\nNo matches."
            total = sum(l["debit"] for l in lines)
            result = [f"=== BILL LINES {keywords} - {len(lines)} lines, TOTAL: ${total:.2f} ==="]
            for l in lines:
                result.append(
                    f"  {l['move_id'][1] if l['move_id'] else '-'} | {l.get('date') or '-'} | "
                    f"{l['name'][:50]} | product: {l['product_id'][1] if l.get('product_id') else '-'} | "
                    f"account: {l['account_id'][1] if l.get('account_id') else '-'} | "
                    f"debit: ${l['debit']:.2f} | vendor: {l['partner_id'][1] if l.get('partner_id') else '-'}")
            return "\n".join(result)
        except Exception as ex:
            return f"[Line search error: {ex}]"

    def _ctx_fleet(self):
        try:
            if "fleet.vehicle" not in self.env:
                return "=== FLEET ===\nFleet module not installed."
            vehicles = self.env["fleet.vehicle"].sudo().search_read(
                [],
                ["id", "name", "license_plate", "model_id", "driver_id",
                 "state_id", "odometer", "acquisition_date",
                 "x_reefer", "x_liftgate", "x_air_ride",
                 "x_max_pallets", "x_max_payload_lbs", "x_gvwr_lbs",
                 "x_vehicle_height_ft", "x_overall_length_ft", "x_cargo_box_length_ft",
                 "x_box_interior_height_ft", "x_door_clearance_in", "x_dock_height_in",
                 "x_tank_capacity_l", "x_def_tank_capacity_l", "x_reserve_fuel_percent",
                 "x_current_fuel_percent", "x_estimated_range_km",
                 "x_avg_km_per_l_last_week", "x_last_location_address", "x_home_base_address"],
                limit=30,
            )
            if not vehicles:
                return "=== FLEET ===\nNo vehicles."
            lines = [f"=== FLEET - {len(vehicles)} vehicles ==="]
            for v in vehicles:
                caps = []
                if v.get("x_reefer"):   caps.append("reefer")
                if v.get("x_liftgate"): caps.append("liftgate")
                if v.get("x_air_ride"): caps.append("air-ride")
                lines.append(
                    f"  {v['name']} | plate: {v.get('license_plate') or '-'} | "
                    f"model: {v['model_id'][1] if v.get('model_id') else '-'} | "
                    f"driver: {v['driver_id'][1] if v.get('driver_id') else '-'} | "
                    f"odometer: {v.get('odometer', 0):.0f}km | "
                    f"state: {v['state_id'][1] if v.get('state_id') else '-'}\n"
                    f"    caps: {', '.join(caps) or 'none'} | "
                    f"payload={v.get('x_max_payload_lbs', 0):.0f}lbs "
                    f"gvwr={v.get('x_gvwr_lbs', 0):.0f}lbs "
                    f"pallets={v.get('x_max_pallets', 0)}\n"
                    f"    dims: h={v.get('x_vehicle_height_ft', 0):.1f}ft "
                    f"len={v.get('x_overall_length_ft', 0):.1f}ft "
                    f"box={v.get('x_cargo_box_length_ft', 0):.1f}ft "
                    f"door={v.get('x_door_clearance_in', 0):.0f}in "
                    f"dock={v.get('x_dock_height_in', 0):.0f}in\n"
                    f"    tank={v.get('x_tank_capacity_l', 0):.0f}L "
                    f"fuel={v.get('x_current_fuel_percent', 0):.0f}% "
                    f"range={v.get('x_estimated_range_km', 0):.0f}km "
                    f"eff={v.get('x_avg_km_per_l_last_week', 0):.2f}km/L\n"
                    f"    loc: {v.get('x_last_location_address') or 'unknown'} | "
                    f"home: {v.get('x_home_base_address') or 'not set'}"
                )
            # Fuel logs
            try:
                fuels = self.env["fleet.vehicle.log.fuel"].sudo().search_read(
                    [], ["vehicle_id", "date", "liter", "price_per_liter", "amount"],
                    order="date desc", limit=20)
                if fuels:
                    total_fuel = sum(f.get("amount", 0) for f in fuels)
                    lines.append(f"\n=== FUEL LOGS (recent 20) - total: ${total_fuel:.2f} ===")
                    for f in fuels:
                        lines.append(
                            f"  {f['vehicle_id'][1] if f.get('vehicle_id') else '-'} | "
                            f"{f.get('date') or '-'} | {f.get('liter',0)}L | "
                            f"${f.get('amount',0):.2f}")
            except Exception:
                pass
            # Service logs
            try:
                svcs = self.env["fleet.vehicle.log.services"].sudo().search_read(
                    [], ["vehicle_id", "date", "service_type_id", "amount"],
                    order="date desc", limit=20)
                if svcs:
                    total_svc = sum(s.get("amount", 0) for s in svcs)
                    lines.append(f"\n=== SERVICE LOGS (recent 20) - total: ${total_svc:.2f} ===")
                    for s in svcs:
                        lines.append(
                            f"  {s['vehicle_id'][1] if s.get('vehicle_id') else '-'} | "
                            f"{s.get('date') or '-'} | "
                            f"{s['service_type_id'][1] if s.get('service_type_id') else '-'} | "
                            f"${s.get('amount',0):.2f}")
            except Exception:
                pass
            return "\n".join(lines)
        except Exception as ex:
            return f"[Fleet error: {ex}]"

    def _ctx_payments(self, pn=None):
        try:
            dom = []
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            pays = self.env["account.payment"].sudo().search_read(
                dom, ["id", "name", "date", "amount", "state",
                       "payment_type", "partner_id", "journal_id"],
                order="date desc", limit=20)
            if not pays:
                return "=== PAYMENTS ===\nNone."
            lines = [f"=== PAYMENTS - {len(pays)} ==="]
            for p in pays:
                lines.append(
                    f"  {p['name']} | {p.get('date')} | ${p['amount']:.2f} | "
                    f"{p.get('payment_type')} | {p['state']} | "
                    f"{p['partner_id'][1] if p.get('partner_id') else '-'}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Payments error: {ex}]"

    def _ctx_subscriptions(self, pn=None):
        try:
            dom = [("is_subscription", "=", True)]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            subs = self.env["sale.order"].sudo().search_read(
                dom, ["id", "name", "partner_id", "state", "amount_total",
                       "next_invoice_date", "recurring_monthly", "subscription_state"],
                limit=15)
            if not subs:
                return f"=== SUBSCRIPTIONS ===\nNone."
            lines = [f"=== SUBSCRIPTIONS - {len(subs)} ==="]
            for s in subs:
                lines.append(
                    f"  {s['name']} | {s['partner_id'][1] if s['partner_id'] else '-'} | "
                    f"{s.get('subscription_state') or s.get('state')} | ${s['amount_total']:.2f}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Subscriptions error: {ex}]"

    def _ctx_crm(self):
        try:
            leads = self.env["crm.lead"].sudo().search_read(
                [("active", "=", True)],
                ["id", "name", "partner_name", "user_id", "stage_id",
                 "expected_revenue", "probability", "city", "phone", "email_from"],
                order="write_date desc", limit=20)
            if not leads:
                return "=== CRM ===\nNo leads."
            lines = [f"=== CRM LEADS - {len(leads)} ==="]
            for l in leads:
                lines.append(
                    f"  {l['name']} | {l.get('partner_name') or '-'} | "
                    f"stage: {l['stage_id'][1] if l.get('stage_id') else '-'} | "
                    f"city: {l.get('city') or '-'} | ${l.get('expected_revenue',0):.0f}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[CRM error: {ex}]"

    def _ctx_tasks(self):
        try:
            tasks = self.env["project.task"].sudo().search_read(
                [("state", "not in", ["1_done", "1_canceled"])],
                ["id", "name", "project_id", "date_deadline", "stage_id"],
                order="date_deadline asc", limit=20)
            if not tasks:
                return "=== TASKS ===\nNone."
            lines = [f"=== TASKS - {len(tasks)} ==="]
            for t in tasks:
                lines.append(
                    f"  {t['name']} | {t['project_id'][1] if t.get('project_id') else '-'} | "
                    f"deadline: {t.get('date_deadline') or '-'}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Tasks error: {ex}]"

    def _ctx_expenses(self):
        try:
            exps = self.env["hr.expense"].sudo().search_read(
                [("state", "in", ["draft", "reported", "approved"])],
                ["id", "name", "employee_id", "total_amount", "date", "state"],
                limit=25)
            if not exps:
                return "=== EXPENSES ===\nNone."
            lines = [f"=== EXPENSES - {len(exps)} ==="]
            for ex in exps:
                lines.append(
                    f"  {ex['name']} | {ex['employee_id'][1] if ex.get('employee_id') else '-'} | "
                    f"${ex['total_amount']:.2f} | {ex.get('date')} | {ex['state']}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Expenses error: {ex}]"

    def _ctx_sales(self):
        try:
            ords = self.env["sale.order"].sudo().search_read(
                [("state", "in", ["sale", "done"])],
                ["id", "name", "partner_id", "amount_total", "date_order", "invoice_status"],
                order="date_order desc", limit=20)
            if not ords:
                return "=== SALES ===\nNone."
            lines = [f"=== SALES - {len(ords)} ==="]
            for o in ords:
                lines.append(
                    f"  {o['name']} | {o['partner_id'][1] if o.get('partner_id') else '-'} | "
                    f"${o['amount_total']:.2f} | {o.get('date_order')}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Sales error: {ex}]"

    def _ctx_journal_entries(self, pn=None):
        try:
            dom = [("state", "=", "posted")]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            entries = self.env["account.move"].sudo().search_read(
                dom, ["id", "name", "date", "partner_id", "journal_id",
                       "amount_total", "move_type"],
                order="date desc", limit=20)
            if not entries:
                return "=== JOURNAL ENTRIES ===\nNone."
            lines = [f"=== JOURNAL ENTRIES - {len(entries)} ==="]
            for e in entries:
                lines.append(
                    f"  {e['name']} | {e.get('date')} | "
                    f"{e['partner_id'][1] if e.get('partner_id') else '-'} | "
                    f"{e['journal_id'][1] if e.get('journal_id') else '-'} | ${e['amount_total']:.2f}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Journal error: {ex}]"

    def _ctx_documents(self, pn=None):
        try:
            dom = [("res_model", "in", ["account.move", "sale.order", "purchase.order"])]
            if pn:
                dom.append(("name", "ilike", pn))
            docs = self.env["ir.attachment"].sudo().search_read(
                dom, ["id", "name", "res_model", "res_id", "mimetype", "file_size"],
                order="create_date desc", limit=20)
            if not docs:
                return "=== DOCUMENTS ===\nNone."
            lines = [f"=== DOCUMENTS - {len(docs)} ==="]
            for d in docs:
                lines.append(
                    f"  {d['name']} | {d.get('res_model')}:{d.get('res_id')} | "
                    f"{round((d.get('file_size') or 0)/1024,1)}KB")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Documents error: {ex}]"

    def _ctx_purchases(self, pn=None):
        try:
            dom = [("state", "in", ["purchase", "done"])]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            pos = self.env["purchase.order"].sudo().search_read(
                dom, ["id", "name", "partner_id", "amount_total", "date_order", "state"],
                order="date_order desc", limit=20)
            if not pos:
                return "=== PURCHASE ORDERS ===\nNone."
            lines = [f"=== PURCHASE ORDERS - {len(pos)} ==="]
            for p in pos:
                lines.append(
                    f"  {p['name']} | {p['partner_id'][1] if p.get('partner_id') else '-'} | "
                    f"${p['amount_total']:.2f} | {p.get('date_order')}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Purchase error: {ex}]"

    def _ctx_accounting_health(self):
        try:
            e = self.env
            return (
                "=== ACCOUNTING HEALTH ===\n"
                f"Unpaid AR: {e['account.move'].sudo().search_count([('move_type','=','out_invoice'),('state','=','posted'),('payment_state','in',['not_paid','partial'])])}\n"
                f"Unpaid AP: {e['account.move'].sudo().search_count([('move_type','=','in_invoice'),('state','=','posted'),('payment_state','in',['not_paid','partial'])])}\n"
                f"Draft bills: {e['account.move'].sudo().search_count([('move_type','=','in_invoice'),('state','=','draft')])}\n"
                f"Bills missing ref: {e['account.move'].sudo().search_count([('move_type','=','in_invoice'),('state','=','posted'),('ref','=',False)])}\n"
                f"Bills >$100 no tax: {e['account.move'].sudo().search_count([('move_type','=','in_invoice'),('state','=','posted'),('amount_tax','=',0),('amount_total','>',100)])}")
        except Exception as ex:
            return f"[Health error: {ex}]"

    def _ctx_partner_full_audit(self, pn):
        if not pn:
            return ""
        parts = []
        try:
            moves = self.env["account.move"].sudo().search_read(
                [("partner_id.name", "ilike", pn), ("state", "!=", "cancel")],
                ["id", "name", "move_type", "state", "invoice_date",
                 "amount_total", "amount_residual", "amount_tax", "payment_state", "ref"],
                order="invoice_date desc", limit=30)
            if moves:
                tot = sum(m["amount_total"] for m in moves)
                owing = sum(m["amount_residual"] for m in moves)
                lines = [f"=== AUDIT '{pn}' - {len(moves)} moves ==="]
                for m in moves:
                    lines.append(
                        f"  {m['name']} | {m['move_type']} | {m['state']} | "
                        f"{m.get('invoice_date') or '-'} | ref: {m.get('ref') or '-'} | "
                        f"${m['amount_total']:.2f} | tax: ${m.get('amount_tax',0):.2f} | "
                        f"owing: ${m['amount_residual']:.2f}")
                lines.append(f"\nTOTALS: billed=${tot:.2f} outstanding=${owing:.2f}")
                parts.append("\n".join(lines))

            bill_lines = self.env["account.move.line"].sudo().search_read(
                [("partner_id.name", "ilike", pn),
                 ("move_id.move_type", "=", "in_invoice"),
                 ("move_id.state", "=", "posted"),
                 ("display_type", "not in", ["line_section", "line_note"])],
                ["move_id", "name", "product_id", "account_id", "quantity", "price_unit", "debit"],
                order="date desc", limit=50)
            if bill_lines:
                lines = [f"=== BILL LINES '{pn}' - {len(bill_lines)} ==="]
                for l in bill_lines:
                    lines.append(
                        f"  {l['move_id'][1] if l['move_id'] else '-'} | {l['name'][:50]} | "
                        f"product: {l['product_id'][1] if l.get('product_id') else '-'} | "
                        f"qty: {l.get('quantity',0)} x ${l.get('price_unit',0):.2f} = ${l['debit']:.2f}")
                parts.append("\n".join(lines))
        except Exception as ex:
            parts.append(f"[Audit error: {ex}]")
        return "\n\n".join(parts)

    def _extract_attachment_content_from_record(self, att_record):
        """Extract readable text from a single stored ir.attachment record.
        Tries PDF text layer → PDF OCR → image OCR → OpenAI Vision.
        Returns up to 3000 chars, or empty string on failure.
        """
        try:
            raw_b64 = att_record.datas
            if not raw_b64:
                return ""
            import base64, io as _io
            file_bytes = base64.b64decode(raw_b64)
            mimetype = (att_record.mimetype or "").lower()
            name = (att_record.name or "").lower()

            is_pdf = "pdf" in mimetype or name.endswith(".pdf")
            is_image = mimetype.startswith("image/") or name.endswith(
                (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"))

            if is_pdf:
                # 1) Text layer (fast, works on digital PDFs)
                try:
                    import pdfplumber
                    parts = []
                    with pdfplumber.open(_io.BytesIO(file_bytes)) as pdf:
                        for page in pdf.pages[:6]:
                            t = page.extract_text(layout=True) or ""
                            if t.strip():
                                parts.append(t)
                    text = "\n".join(parts).strip()
                    if text:
                        return text[:3000]
                except Exception:
                    pass
                # 2) OCR fallback (scanned/image PDF)
                try:
                    from pdf2image import convert_from_bytes
                    import pytesseract
                    imgs = convert_from_bytes(file_bytes, dpi=200, first_page=1, last_page=4)
                    texts = [pytesseract.image_to_string(img) for img in imgs]
                    text = "\n".join(t for t in texts if t).strip()
                    if text:
                        return text[:3000]
                except Exception:
                    pass

            if is_image:
                # 3) Tesseract OCR
                try:
                    from PIL import Image
                    import pytesseract
                    img = Image.open(_io.BytesIO(file_bytes))
                    text = pytesseract.image_to_string(img).strip()
                    if text:
                        return text[:2000]
                except Exception:
                    pass
                # 4) OpenAI Vision fallback
                try:
                    api_key = self._get_api_key()
                    if api_key:
                        b64_str = raw_b64 if isinstance(raw_b64, str) else raw_b64.decode()
                        resp = requests.post(
                            "https://api.openai.com/v1/responses",
                            headers={"Authorization": f"Bearer {api_key}",
                                     "Content-Type": "application/json"},
                            json={
                                "model": self._resolve_model("vision"),
                                "instructions": "Extract all visible text from this document image. Return only the raw text.",
                                "input": [{"role": "user", "content": [
                                    {"type": "input_text",
                                     "text": "Extract all text from this document."},
                                    {"type": "input_image",
                                     "image_url": f"data:{att_record.mimetype};base64,{b64_str}"},
                                ]}],
                                "max_output_tokens": 1000,
                                "store": False,
                            },
                            timeout=30,
                        )
                        if resp.status_code == 200:
                            data = resp.json()
                            text = data.get("output_text", "")
                            if not text:
                                for item in data.get("output", []):
                                    if item.get("type") == "message":
                                        for block in item.get("content", []):
                                            if block.get("type") == "output_text":
                                                text = block.get("text", "")
                                                break
                                    if text:
                                        break
                            if text:
                                return text.strip()[:2000]
                except Exception:
                    pass

            return ""
        except Exception as e:
            _logger.debug("Attachment content extraction failed (ID %s): %s",
                          getattr(att_record, "id", "?"), e)
            return ""

    def _ctx_invoice_attachments(self, pn=None, limit=30):
        """Fetch invoice attachments and extract their text content for AI analysis.
        Returns extracted text from up to `limit` most-recent invoice attachments.
        """
        try:
            dom = [("res_model", "=", "account.move"), ("type", "=", "binary")]
            att_metas = self.env["ir.attachment"].sudo().search_read(
                dom,
                ["id", "name", "res_id", "mimetype", "file_size"],
                order="create_date desc",
                limit=limit,
            )
            if not att_metas:
                return "=== INVOICE ATTACHMENTS ===\nNo attachments found on invoice records."

            # Pre-fetch invoice info in one query
            res_ids = list({m["res_id"] for m in att_metas if m.get("res_id")})
            move_map = {}
            if res_ids:
                moves = self.env["account.move"].sudo().search_read(
                    [("id", "in", res_ids)],
                    ["id", "name", "partner_id", "invoice_date", "move_type", "ref"],
                )
                move_map = {m["id"]: m for m in moves}

            # Filter by partner name if requested
            if pn:
                att_metas = [
                    m for m in att_metas
                    if pn.lower() in (
                        (move_map.get(m.get("res_id"), {}).get("partner_id") or [None, ""])[1] or ""
                    ).lower()
                ] or att_metas  # fall back to all if filter returns nothing

            lines = [f"=== INVOICE ATTACHMENTS — {len(att_metas)} files scanned ==="]

            for meta in att_metas:
                file_size = meta.get("file_size") or 0
                move_data = move_map.get(meta.get("res_id"), {})
                invoice_ref = move_data.get("name") or f"Move:{meta.get('res_id','?')}"
                partner = (move_data.get("partner_id") or [None, "-"])[1]
                inv_date = str(move_data.get("invoice_date") or "-")
                move_type = move_data.get("move_type") or "-"
                ext_ref = move_data.get("ref") or "-"

                lines.append(
                    f"\n--- File: {meta['name']} | Invoice: {invoice_ref} | "
                    f"Ref: {ext_ref} | Partner: {partner} | Date: {inv_date} | Type: {move_type} ---"
                )

                if file_size > 10_000_000:
                    lines.append(f"  [Skipped — file too large ({round(file_size/1024/1024,1)} MB)]")
                    continue

                att_obj = self.env["ir.attachment"].sudo().browse(meta["id"])
                text = self._extract_attachment_content_from_record(att_obj)

                if text:
                    lines.append(f"  EXTRACTED TEXT:\n{text[:2500]}")
                else:
                    lines.append(f"  [No extractable text — mimetype: {meta.get('mimetype','?')}]")

            return "\n".join(lines)
        except Exception as ex:
            return f"[Invoice attachments error: {ex}]"

    def _ctx_all_model_attachments(self, pn=None, limit=40):
        """Read attachments from ALL Odoo models (sales, purchases, CRM, partners, tasks, fleet, etc.)
        Extracts text/content from each file so the AI can answer questions about any stored document.
        """
        try:
            # Models to scan — exclude account.move (already done by _ctx_invoice_attachments)
            target_models = [
                "sale.order", "purchase.order", "crm.lead",
                "res.partner", "project.task", "fleet.vehicle",
                "stock.picking", "hr.employee", "mail.message",
            ]
            available = [m for m in target_models if m in self.env]
            if not available:
                return ""

            dom = [
                ("res_model", "in", available),
                ("type", "=", "binary"),
            ]
            if pn:
                dom.append(("name", "ilike", pn))

            att_metas = self.env["ir.attachment"].sudo().search_read(
                dom,
                ["id", "name", "res_model", "res_id", "mimetype", "file_size", "create_date"],
                order="create_date desc",
                limit=limit,
            )
            if not att_metas:
                return ""

            # Batch-fetch record display names for context
            record_map = {}
            for model in available:
                ids_for_model = [m["res_id"] for m in att_metas if m["res_model"] == model and m.get("res_id")]
                if not ids_for_model:
                    continue
                try:
                    recs = self.env[model].sudo().browse(ids_for_model)
                    for r in recs:
                        display = getattr(r, "name", None) or getattr(r, "partner_name", None) or str(r.id)
                        record_map[(model, r.id)] = display
                except Exception:
                    pass

            lines = [f"=== ODOO ATTACHMENTS (all models) — {len(att_metas)} files ==="]

            for meta in att_metas:
                model = meta.get("res_model", "?")
                res_id = meta.get("res_id", "?")
                record_label = record_map.get((model, res_id), f"ID {res_id}")
                file_size = meta.get("file_size") or 0

                lines.append(
                    f"\n--- {meta['name']} | Model: {model} | Record: {record_label} "
                    f"| Type: {meta.get('mimetype','?')} ---"
                )

                if file_size > 10_000_000:
                    lines.append(f"  [Skipped — too large ({round(file_size/1024/1024,1)} MB)]")
                    continue

                att_obj = self.env["ir.attachment"].sudo().browse(meta["id"])
                text = self._extract_attachment_content_from_record(att_obj)
                if text:
                    lines.append(f"  CONTENT:\n{text[:2500]}")
                else:
                    lines.append(f"  [No extractable text]")

            return "\n".join(lines)
        except Exception as ex:
            return f"[All-model attachments error: {ex}]"

    def _ctx_employees(self):
        """List HR employees."""
        try:
            if "hr.employee" not in self.env:
                return ""
            emps = self.env["hr.employee"].sudo().search_read(
                [("active", "=", True)],
                ["id", "name", "job_id", "department_id", "work_phone", "work_email"],
                limit=50,
            )
            if not emps:
                return "=== EMPLOYEES ===\nNone."
            lines = [f"=== EMPLOYEES — {len(emps)} ==="]
            for e in emps:
                lines.append(
                    f"  {e['name']} | job: {e['job_id'][1] if e.get('job_id') else '-'} | "
                    f"dept: {e['department_id'][1] if e.get('department_id') else '-'} | "
                    f"phone: {e.get('work_phone') or '-'} | email: {e.get('work_email') or '-'}"
                )
            return "\n".join(lines)
        except Exception as ex:
            return f"[Employees error: {ex}]"

    def _ctx_stock_pickings(self, pn=None):
        """List recent stock pickings (delivery orders, receipts)."""
        try:
            if "stock.picking" not in self.env:
                return ""
            dom = [("state", "in", ["done", "assigned", "waiting", "confirmed"])]
            if pn:
                dom.append(("partner_id.name", "ilike", pn))
            picks = self.env["stock.picking"].sudo().search_read(
                dom,
                ["id", "name", "partner_id", "picking_type_id", "scheduled_date",
                 "date_done", "state", "origin"],
                order="scheduled_date desc",
                limit=25,
            )
            if not picks:
                return "=== STOCK PICKINGS ===\nNone."
            lines = [f"=== STOCK PICKINGS — {len(picks)} ==="]
            for p in picks:
                lines.append(
                    f"  {p['name']} | {p['picking_type_id'][1] if p.get('picking_type_id') else '-'} | "
                    f"partner: {p['partner_id'][1] if p.get('partner_id') else '-'} | "
                    f"origin: {p.get('origin') or '-'} | state: {p['state']} | "
                    f"done: {p.get('date_done') or p.get('scheduled_date') or '-'}"
                )
            return "\n".join(lines)
        except Exception as ex:
            return f"[Stock pickings error: {ex}]"

    def _ctx_purchase_lines(self, pn=None):
        """List recent purchase order lines for deep cost analysis."""
        try:
            if "purchase.order.line" not in self.env:
                return ""
            dom = [("order_id.state", "in", ["purchase", "done"])]
            if pn:
                dom.append(("order_id.partner_id.name", "ilike", pn))
            lines_data = self.env["purchase.order.line"].sudo().search_read(
                dom,
                ["id", "order_id", "product_id", "name", "product_qty",
                 "price_unit", "price_subtotal", "date_planned"],
                order="date_planned desc",
                limit=40,
            )
            if not lines_data:
                return ""
            lines = [f"=== PURCHASE ORDER LINES — {len(lines_data)} ==="]
            for l in lines_data:
                lines.append(
                    f"  {l['order_id'][1] if l.get('order_id') else '-'} | "
                    f"{l['product_id'][1] if l.get('product_id') else l.get('name','?')[:40]} | "
                    f"qty: {l.get('product_qty',0)} x ${l.get('price_unit',0):.2f} = ${l.get('price_subtotal',0):.2f}"
                )
            return "\n".join(lines)
        except Exception as ex:
            return f"[Purchase lines error: {ex}]"

    def _ctx_account_balances(self):
        """Summary of key account balances."""
        try:
            accounts = self.env["account.account"].sudo().search_read(
                [("account_type", "in", [
                    "asset_receivable", "liability_payable",
                    "income", "expense", "asset_cash", "liability_credit_card",
                ])],
                ["id", "name", "code", "account_type", "current_balance"],
                limit=40,
            )
            if not accounts:
                return ""
            lines = ["=== KEY ACCOUNT BALANCES ==="]
            for a in accounts:
                lines.append(
                    f"  [{a.get('code','')}] {a['name']} | type: {a.get('account_type','-')} | "
                    f"balance: ${a.get('current_balance',0):.2f}"
                )
            return "\n".join(lines)
        except Exception as ex:
            return f"[Account balances error: {ex}]"

    def _ctx_wa_negotiations(self, entity=None):
        """Surface recent WA negotiations for AI context."""
        try:
            domain = [('status', 'in', ('draft', 'negotiating', 'agreed', 'quotation_created'))]
            if entity:
                domain += [('partner_id.name', 'ilike', entity)]
            negs = self.env['premafirm.wa.negotiation'].sudo().search(
                domain, order='create_date desc', limit=10)
            if not negs:
                return ""
            lines = ["=== WA NEGOTIATIONS ==="]
            for n in negs:
                partner = n.partner_id.name if n.partner_id else 'Unknown'
                status_label = dict(n._fields['status'].selection).get(n.status, n.status)
                line = (
                    f"• {partner} | Status: {status_label} | "
                    f"Their Offer: ${n.their_offer:.2f} | "
                    + (f"Our Counter: ${n.our_counter:.2f} | " if n.our_counter else "")
                    + (f"Agreed: ${n.agreed_rate:.2f} CAD | " if n.agreed_rate else "")
                    + f"Stops: {n.stops_count} | {n.commodity or 'freight'} | {n.equipment_type or ''}"
                    + (f" → Quotation: {n.sale_order_id.name}" if n.sale_order_id else "")
                )
                lines.append(line)
            return "\n".join(lines)
        except Exception:
            return ""

    def _ctx_source_library(self):
        """Extract text from all enabled source library files."""
        try:
            cfg = self.env["prema.ai.settings"].get_singleton()
            if not cfg.source_library_enabled or not cfg.source_attachment_ids:
                return ""
            ref_lines = ["=== SOURCE LIBRARY FILES ==="]
            for att in cfg.source_attachment_ids[:10]:
                text = self._extract_attachment_content_from_record(att)
                if text:
                    ref_lines.append(f"\n--- {att.name} ---\n{text[:3000]}")
                else:
                    ref_lines.append(f"\n--- {att.name} --- [binary/unsupported format]")
            return "\n".join(ref_lines)
        except Exception as ex:
            return f"[Source library error: {ex}]"

    def _ctx_pricing_context(self):
        """Load pricing engine history and rules for Pricing Review mode."""
        try:
            lines = ["=== PRICING ENGINE CONTEXT ==="]
            # Historical pricing logs from premafirm_ai_engine
            if "premafirm.ai.log" in self.env:
                logs = self.env["premafirm.ai.log"].sudo().search_read(
                    [], ["lead_id", "distance_km", "pallets", "final_rate"],
                    order="id desc", limit=25)
                if logs:
                    lines.append(f"Recent Pricing History ({len(logs)} records):")
                    for log in logs:
                        lines.append(
                            f"  Lead: {log.get('lead_id', '?')} | "
                            f"{log.get('distance_km', 0)} km | "
                            f"{log.get('pallets', 0)} pallets | "
                            f"Rate: ${log.get('final_rate', 0):.2f}"
                        )
            # Structured pricing rules from config
            cfg = self.env["prema.ai.settings"].get_singleton()
            if cfg.pricing_rules and cfg.pricing_rules.strip():
                lines.append(f"\nPricing Rules:\n{cfg.pricing_rules.strip()}")
            return "\n".join(lines)
        except Exception as ex:
            return f"[Pricing context error: {ex}]"

    # ====================================================================
    # System prompt
    # ====================================================================
    def _build_system_prompt(self, chat_mode="standard"):
        memory = ""
        try:
            memory = self.env["prema.ai.correction"].build_memory_prompt()
        except Exception:
            pass

        custom_instructions = ""
        try:
            cfg = self.env["prema.ai.settings"].get_singleton()
            if cfg.custom_instructions and cfg.custom_instructions.strip():
                custom_instructions = cfg.custom_instructions.strip()
        except Exception:
            pass

        prompt = """You are Prema AI, the CEO-level intelligent assistant for PremaFirm Inc (Canadian trucking & logistics).

CRITICAL: You have FULL READ ACCESS to the entire Odoo database. Live data is provided below. NEVER say "I don't have access" or "I can't check". The data IS there — analyze it.

CAPABILITIES:
- Full read access to EVERY Odoo app: Accounting, CRM, Fleet, Sales, Purchases, Projects, HR, Documents, Stock, Employees
- Full attachment access: You CAN read and extract text/content from ANY file attached to ANY Odoo record —
  invoices, sales orders, purchase orders, CRM leads, partners, tasks, fleet vehicles, and more
- Files include: PDFs, scanned images, photos, Word docs, Excel sheets, signed contracts, PODs, BOLs
- Customer balances, payment history, aging reports — all directly accessible from live ERP data
- Web search: You can search the internet for research, prospects, market info
- File generation: You can create PDF, Excel (.xlsx), Word (.docx), CSV files
- CRM leads: The backend AUTOMATICALLY creates CRM leads when the user asks
- Voice input: Users can speak to you via microphone

CRM LEAD CREATION — CRITICAL RULE:
NEVER say you cannot create CRM records, cannot add to CRM, or lack write access.
The system backend handles ALL record creation automatically.
When user asks to add companies or leads to CRM (in any wording), respond with:
  "I've prepared X lead(s) for CRM. Click the [Create Leads] button above to add them."
Do NOT explain that you are an AI that can't write to databases.
Do NOT say "you would need to manually add these".
The backend intercepts these requests and handles everything.

FILE GENERATION:
Supported: PDF, Excel (.xlsx), Word (.docx), CSV
NOT supported: PowerPoint, images, or other formats
When asked for unsupported format, tell the user what you CAN create.

RULES:
- Read queries: answer immediately with real data
- Record changes: show preview, ask approval
- NEVER delete records
- Bills stay DRAFT until user confirms

READING STORED ATTACHMENTS (ALL ODOO MODELS):
You CAN directly access and analyze files attached to ANY record in Odoo.
When asked about documents, photos, images, contracts, or any file:
- The ERP context below includes EXTRACTED CONTENT from those attachment files
- This covers: invoices, sales orders, purchase orders, CRM leads, partners, tasks, fleet vehicles
- Read the content and identify names, dates, amounts, ref numbers, signatures, or any relevant data
- If asked about a specific record (e.g. "invoice INV/2024/001"), find its attachments in the context
- If a file shows "[No extractable text]" note that file could not be read
- For customer balance questions: use the live account balance and invoice data in the ERP context
- NEVER say "I cannot access attachments" or "I can't see documents" — all content is in the context below. Use it.

UPLOADING NEW FILES TO CHAT:
When a user uploads a file, its content is extracted and included directly in their message as text.
- PDFs: text is extracted from the PDF layer or via OCR
- Excel / CSV: sheet data is included as tab-separated text
- Word documents: paragraph and table text is extracted
- Images: sent directly for visual analysis
Analyze the extracted content immediately and answer the user's question.
If the extracted content says it "could not be extracted", tell the user and suggest pasting the text manually.
Never auto-create records from newly uploaded attachments without user confirmation.

CRM LEAD DUPLICATE PREVENTION:
Before quoting lead counts, the system automatically checks for existing contacts.
If a company already exists in CRM or Contacts, it will NOT be duplicated — the system skips it and notifies you.
Always confirm the final count after duplicate filtering.

RESPONSE STYLE:
- Use actual record names, IDs, amounts, dates from the ERP context
- Format currency as $X.XX
- Use clean Markdown: **bold** for emphasis, ## for section headers, - for bullets, 1. for steps
- NEVER output raw HTML tags like &lt;div&gt;, &lt;br&gt;, &lt;p&gt;, &lt;strong&gt; etc.
- When you find issues, explain what's wrong and how to fix it"""
        # ── Source priority hierarchy ──────────────────────────────────
        prompt += (
            "\n\nSOURCE PRIORITY — ALWAYS FOLLOW THIS ORDER:\n"
            "1. System instructions and company rules (this prompt)\n"
            "2. Internal Odoo database — live records, balances, transactions, ALL apps\n"
            "3. Files and attachments stored in Odoo — documents, images, scans on any record\n"
            "4. Source library files (rate cards, SOPs, agreements uploaded in AI Settings)\n"
            "5. Connected APIs (pricing engine, dispatch rules)\n"
            "6. Web research (LAST RESORT — only when none of the above have the answer)\n"
            "NEVER use web search if the answer is available in the Odoo database or attached files."
        )

        # ── Mode-specific instructions ─────────────────────────────────
        _mode_instructions = {
            "deep_thinking": (
                "\n\nCURRENT MODE: DEEP THINKING\n"
                "- Use multi-step reasoning before answering\n"
                "- Break complex problems into clearly numbered steps\n"
                "- Show your analysis process before stating conclusions\n"
                "- Cross-reference multiple data sources and double-check calculations\n"
                "- Provide thorough, well-structured analysis even if the response is longer"
            ),
            "web_research": (
                "\n\nCURRENT MODE: WEB RESEARCH\n"
                "- Actively use web search for current information\n"
                "- Validate and supplement internal data with external sources\n"
                "- Report relevant news, market rates, and industry trends\n"
                "- Always cite the source of web-researched information\n"
                "- Prefer real-time external data when the question is about current events"
            ),
            "pricing_review": (
                "\n\nCURRENT MODE: PRICING REVIEW\n"
                "- Focus exclusively on rate calculation and profitability analysis\n"
                "- Apply company pricing rules and rate tables strictly\n"
                "- Compare proposed rates against historical pricing data in context\n"
                "- Flag loads below minimum margin thresholds\n"
                "- Always include fuel surcharge, pallet fees, and applicable accessorials\n"
                "- Do NOT use web search — use only internal pricing data and rules"
            ),
            "lead_generation": (
                "\n\nCURRENT MODE: LEAD GENERATION\n"
                "CRITICAL: NEVER say 'hold on', 'please wait', 'give me a moment', 'I will search' "
                "or any promise of future action. Do ALL work inline and return results immediately.\n"
                "- The system automatically handles web research — do not describe what you are about to do\n"
                "- If the system has already returned lead results above, DO NOT repeat the search — just present the results\n"
                "- You CAN create tags: say 'Create tag [name]' to create CRM and contact tags instantly\n"
                "- You CAN answer questions about the leads already in CRM\n"
                "- You CAN advise on outreach strategy, email templates, and follow-up timing\n"
                "- Prioritise: Logistics Managers, VP Supply Chain, Director of Transportation, "
                "Procurement Managers, Operations Directors at shippers, brokers, and distributors"
            ),
            "document_qa": (
                "\n\nCURRENT MODE: DOCUMENT Q&A\n"
                "- Prioritise answers from stored Odoo documents, files, and attachments\n"
                "- You have FULL access to all Odoo records AND every file/image attached to them\n"
                "- Read and analyse attached documents on invoices, sales orders, purchase orders, "
                "CRM leads, partners, tasks, and fleet vehicles\n"
                "- Quote or cite the specific document name when answering from a file\n"
                "- Use ERP live data (balances, dates, amounts) to supplement document content\n"
                "- If a specific document is not found, say so and offer what the ERP data shows"
            ),
            "logistics_estimate": (
                "\n\nCURRENT MODE: LOGISTICS ESTIMATE\n"
                "- You are a logistics cost and routing intelligence assistant for a Canadian trucking company\n"
                "- LOGISTICS ESTIMATE DATA is provided in the ERP context — use it as your primary source\n"
                "- Present a structured estimate including: selected truck, route summary, truck-safe ETA, "
                "traffic alerts, weather warnings, fuel estimate, refuel recommendation, feasibility result, "
                "cost estimate, and suggested final logistics rate\n"
                "- If LOGISTICS ESTIMATE DATA is present, do NOT re-derive the numbers — use them directly\n"
                "- If the 'truck' key is present and non-empty in LOGISTICS ESTIMATE DATA, the user has ALREADY "
                "selected that truck — do NOT ask them to confirm or re-select the truck\n"
                "- Only ask for the truck if 'truck' is absent or empty in LOGISTICS ESTIMATE DATA\n"
                "- Apply Cost Calculation Rules and Routing / Dispatch Rules from AI Settings\n"
                "- Flag any safety issues: low fuel, severe weather, bridge restrictions, urgent timelines\n"
                "- If truck data is missing or the service was unavailable, note this and provide a best-effort estimate\n"
                "- Format output with clear section headers: ## Truck, ## Route, ## Fuel & Range, "
                "## Refuel Strategy, ## Cost Estimate, ## Recommendation\n"
                "- Currency in CAD. Distances in km. Volume in litres."
            ),
        }
        if chat_mode in _mode_instructions:
            prompt += _mode_instructions[chat_mode]

        # ── General custom instructions ────────────────────────────────
        if custom_instructions:
            prompt += f"\n\nCOMPANY-SPECIFIC INSTRUCTIONS:\n{custom_instructions}"

        # ── Structured rules framework (from AI Settings) ──────────────
        try:
            cfg = self.env["prema.ai.settings"].get_singleton()
            if cfg.pricing_rules and cfg.pricing_rules.strip() and chat_mode in ("standard", "pricing_review", "logistics_estimate"):
                prompt += f"\n\nPRICING ENGINE RULES:\n{cfg.pricing_rules.strip()}"
            if cfg.lead_gen_instructions and cfg.lead_gen_instructions.strip() and chat_mode in ("standard", "lead_generation"):
                prompt += f"\n\nLEAD GENERATION INSTRUCTIONS:\n{cfg.lead_gen_instructions.strip()}"
            if cfg.duplicate_prevention_rules and cfg.duplicate_prevention_rules.strip():
                prompt += f"\n\nDUPLICATE PREVENTION RULES:\n{cfg.duplicate_prevention_rules.strip()}"
            if cfg.routing_dispatch_rules and cfg.routing_dispatch_rules.strip() and chat_mode in ("standard", "pricing_review", "logistics_estimate"):
                prompt += f"\n\nROUTING / DISPATCH RULES:\n{cfg.routing_dispatch_rules.strip()}"
            if cfg.cost_calc_rules and cfg.cost_calc_rules.strip() and chat_mode == "logistics_estimate":
                prompt += f"\n\nCOST CALCULATION RULES:\n{cfg.cost_calc_rules.strip()}"
            if cfg.messaging_rules and cfg.messaging_rules.strip():
                prompt += f"\n\nMESSAGING RULES:\n{cfg.messaging_rules.strip()}"
        except Exception:
            pass

        if memory:
            prompt += f"\n\n{memory}"
        return prompt

    # ====================================================================
    # OpenAI Responses API with web search
    # ====================================================================
    def _call_openai(self, mode="primary", erp_context="", file_contents=None, chat_mode="standard"):
        self.ensure_one()
        api_key = self._get_api_key()
        if not api_key:
            return ("\u26A0 API key not found.\n"
                    "Set it in: Prema AI > Settings > prema_ai.api_key")

        model = self._resolve_model(mode)
        # Only upgrade to vision model if there are actual images (not text-extractable files)
        if file_contents:
            has_images = any(
                (fc.get("mimetype") or "").startswith("image/")
                for fc in file_contents
            )
            if has_images:
                model = self._resolve_model("vision")

        instructions = self._build_system_prompt(chat_mode=chat_mode)
        input_messages = []

        if erp_context:
            input_messages.append({
                "role": "user",
                "content": "[LIVE ERP DATABASE]\n\n" + erp_context,
            })
            input_messages.append({
                "role": "assistant",
                "content": "ERP data loaded. Analyzing now.",
            })

        for msg in self.message_ids:
            role = msg.role if msg.role in ("user", "assistant") else "user"
            content = msg.content or ""
            if content.startswith("[LIVE ERP DATABASE"):
                continue
            input_messages.append({"role": role, "content": content})

        # File content injection — handles images, PDFs, Excel, Word, CSV
        has_images_in_content = any(
            (fc.get("mimetype") or "").startswith("image/")
            for fc in (file_contents or [])
        )
        if file_contents and input_messages:
            last_user = None
            for i in range(len(input_messages) - 1, -1, -1):
                if input_messages[i]["role"] == "user":
                    last_user = i
                    break
            if last_user is not None:
                raw_text = input_messages[last_user]["content"]
                raw_text = raw_text if isinstance(raw_text, str) else str(raw_text)

                # When user pastes an image with no text (or only the auto filename),
                # replace with an explicit instruction so GPT knows what to do.
                is_image_only = (
                    raw_text.startswith("\U0001F4CE")  # 📎 emoji prefix
                    or not raw_text.strip()
                )
                if is_image_only and has_images_in_content:
                    raw_text = (
                        "Please analyze this image carefully. "
                        "Extract ALL visible text exactly as it appears. "
                        "Identify the document type (invoice, BOL, rate sheet, delivery note, screenshot, etc.) "
                        "and summarize every key piece of information you can see."
                    )

                parts = [{"type": "input_text", "text": raw_text}]

                for fc in file_contents:
                    mime = (fc.get("mimetype") or "application/octet-stream").lower()
                    fname = fc.get("filename", "file")
                    b64 = fc.get("b64", "")

                    if mime.startswith("image/"):
                        # Images go directly as input_image (vision)
                        _logger.info(
                            "Vision: sending %s image (%s) to model %s",
                            fname, mime, model
                        )
                        parts.append({
                            "type": "input_image",
                            "image_url": f"data:{mime};base64,{b64}",
                        })
                    else:
                        # All other files: extract text on the Python side first
                        extracted = self._extract_b64_file_content(b64, mime, fname)
                        if extracted:
                            parts.append({
                                "type": "input_text",
                                "text": f"[Content of '{fname}' ({mime})]:\n{extracted}",
                            })
                        else:
                            parts.append({
                                "type": "input_text",
                                "text": (
                                    f"[File '{fname}' was attached but its content could not be "
                                    f"extracted automatically (type: {mime}). "
                                    f"If it is a scanned image, ask the user to paste the text.]"
                                ),
                            })

                input_messages[last_user] = {"role": "user", "content": parts}

        # Mode-specific temperature and token limits
        _temperature_map = {
            "standard": 0.3,
            "deep_thinking": 0.1,
            "web_research": 0.4,
            "pricing_review": 0.1,
            "lead_generation": 0.4,
            "document_qa": 0.0,
        }
        _max_tokens_map = {
            "deep_thinking": 8192,
        }
        temperature = _temperature_map.get(chat_mode, 0.3)
        max_tokens = _max_tokens_map.get(chat_mode, 4096)

        # Web search: disable when images are present (tool routing conflicts with vision
        # in the Responses API — model may skip image analysis when tools are active)
        _no_web_modes = {"pricing_review", "document_qa"}
        tools = [] if (chat_mode in _no_web_modes or has_images_in_content) else [{"type": "web_search_preview"}]

        body = {
            "model": model,
            "instructions": instructions,
            "input": input_messages,
            "max_output_tokens": max_tokens,
            "store": False,
        }
        if self._supports_temperature(model):
            body["temperature"] = temperature
        if tools:
            body["tools"] = tools

        try:
            resp = requests.post(
                "https://api.openai.com/v1/responses",
                headers={"Authorization": f"Bearer {api_key}",
                         "Content-Type": "application/json"},
                json=body, timeout=120)

            if resp.status_code != 200:
                err = ""
                try:
                    err = resp.json().get("error", {}).get("message", resp.text[:300])
                except Exception:
                    err = resp.text[:300]
                return (f"\u26A0 OpenAI API Error (HTTP {resp.status_code}):\n{err}\n\n"
                        f"Model: {model}\nCheck Prema AI > Settings.")

            data = resp.json()
            text = data.get("output_text")
            if text:
                return text.strip()

            for item in data.get("output", []):
                if item.get("type") == "message":
                    for block in item.get("content", []):
                        if block.get("type") == "output_text" and block.get("text"):
                            return block["text"].strip()

            return "\u26A0 AI returned empty response. Try again."

        except requests.exceptions.Timeout:
            return "\u26A0 Request timed out. Try a simpler question."
        except requests.exceptions.ConnectionError as e:
            return f"\u26A0 Cannot reach OpenAI: {e}"
        except Exception as e:
            _logger.error("OpenAI error: %s\n%s", e, traceback.format_exc())
            return f"\u26A0 Error: {e}\nCheck Odoo logs."

    # ── Tool registry ──────────────────────────────────────────────
    def _tool_registry(self):
        return {
            "search_records": lambda m, d, f=None: self.env[m].sudo().search_read(d, f),
            "financial_summary": self._tool_financial_summary,
            "crm_pipeline_analysis": self._tool_crm_pipeline_analysis,
        }

    def _tool_financial_summary(self):
        inc = sum(self.env["account.move"].sudo().search(
            [("move_type", "=", "out_invoice")]).mapped("amount_total"))
        exp = sum(self.env["account.move"].sudo().search(
            [("move_type", "=", "in_invoice")]).mapped("amount_total"))
        return {"income": inc, "expenses": exp, "balance": inc - exp}

    def _tool_crm_pipeline_analysis(self):
        leads = self.env["crm.lead"].sudo().search([])
        sc = {}
        for ld in leads:
            s = ld.stage_id.name or "No Stage"
            sc[s] = sc.get(s, 0) + 1
        return {"stages": sc}
