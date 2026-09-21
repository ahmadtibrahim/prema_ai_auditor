# FILE: /opt/odoo/custum-addons/prema_ai_auditor/services/file_generator.py
"""
File generator: creates PDF, Excel, Word, CSV from structured data.
"""
import csv
import io
import logging
from datetime import datetime

_logger = logging.getLogger(__name__)


def generate_file(fmt, report_data):
    """
    Generate a file in the specified format.
    report_data: {"title": "...", "headers": [...], "rows": [[...], ...], "summary": "..."}
    Returns: (file_bytes, filename)
    """
    title = report_data.get("title", "Prema AI Report")
    headers = report_data.get("headers", [])
    rows = report_data.get("rows", [])
    summary = report_data.get("summary", "")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    safe_title = title.replace(" ", "_")[:30]

    if fmt == "xlsx":
        return _gen_xlsx(title, headers, rows, summary, f"{safe_title}_{timestamp}.xlsx")
    elif fmt == "docx":
        return _gen_docx(title, headers, rows, summary, f"{safe_title}_{timestamp}.docx")
    elif fmt == "csv":
        return _gen_csv(headers, rows, f"{safe_title}_{timestamp}.csv")
    else:
        return _gen_pdf(title, headers, rows, summary, f"{safe_title}_{timestamp}.pdf")


def _gen_xlsx(title, headers, rows, summary, filename):
    buf = io.BytesIO()
    import xlsxwriter
    wb = xlsxwriter.Workbook(buf)
    ws = wb.add_worksheet("Report")

    # Formats
    title_fmt = wb.add_format({"bold": True, "font_size": 16, "bottom": 2})
    header_fmt = wb.add_format({
        "bold": True, "bg_color": "#4f46e5", "font_color": "white",
        "border": 1, "text_wrap": True})
    cell_fmt = wb.add_format({"border": 1, "text_wrap": True})
    money_fmt = wb.add_format({"border": 1, "num_format": "$#,##0.00"})
    summary_fmt = wb.add_format({"italic": True, "font_size": 10, "text_wrap": True})

    # Title
    ws.merge_range(0, 0, 0, max(len(headers) - 1, 1), title, title_fmt)
    ws.write(1, 0, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    # Headers
    row_idx = 3
    for col, h in enumerate(headers):
        ws.write(row_idx, col, h, header_fmt)
        ws.set_column(col, col, max(15, len(str(h)) + 5))

    # Data rows
    for row in rows:
        row_idx += 1
        for col, val in enumerate(row):
            if isinstance(val, (int, float)):
                ws.write_number(row_idx, col, val, money_fmt)
            else:
                ws.write(row_idx, col, str(val) if val else "", cell_fmt)

    # Summary
    if summary:
        row_idx += 2
        ws.merge_range(row_idx, 0, row_idx + 2, max(len(headers) - 1, 3), summary, summary_fmt)

    wb.close()
    return buf.getvalue(), filename


def _gen_pdf(title, headers, rows, summary, filename):
    buf = io.BytesIO()
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    # Use landscape if many columns
    page = landscape(letter) if len(headers) > 5 else letter
    doc = SimpleDocTemplate(buf, pagesize=page)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | PremaFirm Inc.",
        styles["Normal"]))
    elements.append(Spacer(1, 20))

    # Table
    if headers and rows:
        table_data = [headers] + [[str(v) if v else "" for v in row] for row in rows]
        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elements.append(t)

    # Summary
    if summary:
        elements.append(Spacer(1, 20))
        elements.append(Paragraph(summary, styles["Normal"]))

    doc.build(elements)
    return buf.getvalue(), filename


def _gen_docx(title, headers, rows, summary, filename):
    buf = io.BytesIO()
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    t = doc.add_heading(title, level=1)
    doc.add_paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | PremaFirm Inc.")
    doc.add_paragraph("")

    # Table
    if headers and rows:
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Light Grid Accent 1"

        # Headers
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = str(h)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True

        # Data
        for row_idx, row in enumerate(rows):
            for col_idx, val in enumerate(row):
                table.rows[row_idx + 1].cells[col_idx].text = str(val) if val else ""

    # Summary
    if summary:
        doc.add_paragraph("")
        doc.add_paragraph(summary)

    doc.save(buf)
    return buf.getvalue(), filename


def _gen_csv(headers, rows, filename):
    buf = io.StringIO()
    writer = csv.writer(buf)
    if headers:
        writer.writerow(headers)
    for row in rows:
        writer.writerow([str(v) if v else "" for v in row])
    return buf.getvalue().encode("utf-8"), filename
